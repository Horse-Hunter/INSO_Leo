"""Idempotent, atomic local Excel persistence for Research V1."""

from __future__ import annotations

import os
import tempfile
from collections.abc import Mapping
from copy import copy
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Final

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from openpyxl.workbook.workbook import Workbook as OpenpyxlWorkbook

from .source_contracts import ResearchSource

MPN_HEADER = "型号"
BRAND_HEADER = "品牌"
QUANTITY_HEADER = "数量"
IMPORTANCE_HEADER = "重要等级"
STOCK_HEADER = "货量标识"
ESTIMATED_TOTAL_HEADER = "预估订单总价"
MARKET_REFERENCE_HEADER = "市场最低参考价"
INSO_HEADER = "INSO"
FINDCHIPS_HEADER = "Findchips"
HQEW_HEADER = "华强"
LCSC_HEADER = "立创"
BOM_AI_HEADER = "正能量"
REMARKS_HEADER = "备注"
PROCESSED_AT_HEADER = "处理时间"
INQUIRY_ID_HEADER = "_inquiry_id"
RESEARCH_STATUS_HEADER = "_research_status"
VISIBLE_HEADERS = (
    MPN_HEADER,
    BRAND_HEADER,
    QUANTITY_HEADER,
    IMPORTANCE_HEADER,
    STOCK_HEADER,
    ESTIMATED_TOTAL_HEADER,
    MARKET_REFERENCE_HEADER,
    INSO_HEADER,
    FINDCHIPS_HEADER,
    HQEW_HEADER,
    LCSC_HEADER,
    BOM_AI_HEADER,
    REMARKS_HEADER,
    PROCESSED_AT_HEADER,
)
CANONICAL_HEADERS = (*VISIBLE_HEADERS, RESEARCH_STATUS_HEADER, INQUIRY_ID_HEADER)
DEFAULT_SHEET_TITLE = "Research"
SOURCE_HEADERS: Mapping[ResearchSource, str] = {
    ResearchSource.INSO: INSO_HEADER,
    ResearchSource.FINDCHIPS: FINDCHIPS_HEADER,
    ResearchSource.HQEW: HQEW_HEADER,
    ResearchSource.LCSC: LCSC_HEADER,
    ResearchSource.BOM_AI: BOM_AI_HEADER,
}

_OLD_TOTAL_HEADER = "预计订单总价"
_OLD_CANONICAL_HEADERS = (
    MPN_HEADER, BRAND_HEADER, QUANTITY_HEADER, IMPORTANCE_HEADER, STOCK_HEADER,
    ESTIMATED_TOTAL_HEADER, MARKET_REFERENCE_HEADER, INSO_HEADER, FINDCHIPS_HEADER,
    HQEW_HEADER, LCSC_HEADER, BOM_AI_HEADER, REMARKS_HEADER, INQUIRY_ID_HEADER,
)
_PREVIOUS_HEADERS = (
    MPN_HEADER,
    BRAND_HEADER,
    QUANTITY_HEADER,
    IMPORTANCE_HEADER,
    STOCK_HEADER,
    _OLD_TOTAL_HEADER,
    MARKET_REFERENCE_HEADER,
    REMARKS_HEADER,
    INQUIRY_ID_HEADER,
)
_LEGACY_HEADERS = (INQUIRY_ID_HEADER, IMPORTANCE_HEADER, REMARKS_HEADER)


class _UnsetType:
    pass


_UNSET: Final = _UnsetType()


class ExcelOutputError(RuntimeError):
    """Base error for Research-owned Excel output."""


class ExcelConsistencyError(ExcelOutputError):
    """Persisted workbook identity or schema is ambiguous."""


class ExcelWriteError(ExcelOutputError):
    """The workbook could not be atomically saved."""


@dataclass(frozen=True, slots=True)
class ResearchHistoryRecord:
    inquiry_id: str
    mpn: str | None
    brand: str | None
    quantity: object | None
    importance_raw: str | None
    stock_label: str | None
    estimated_total: object | None
    market_reference: str | None
    source_values: tuple[tuple[str, str], ...]
    remarks: str | None
    processed_at: datetime | None
    research_status: str | None


def importance_display_value(importance_raw: str | None) -> str | None:
    """The canonical workbook displays the normalized raw grade unchanged."""

    return importance_raw


def _decimal_text(value: Decimal | None) -> str | None:
    from .source_contracts import money_text

    return None if value is None else money_text(value)


class ResearchExcelOutput:
    """Persist one full logical snapshot per inquiry_id."""

    def __init__(self, path: str | Path) -> None:
        self.path = Path(path)

    def upsert(
        self,
        inquiry_id: str,
        *,
        importance_raw: str | None,
        remarks: str | None | _UnsetType = _UNSET,
        mpn: str | None | _UnsetType = _UNSET,
        brand: str | None | _UnsetType = _UNSET,
        quantity: int | None | _UnsetType = _UNSET,
        stock_label: str | None | _UnsetType = _UNSET,
        estimated_total: Decimal | None | _UnsetType = _UNSET,
        market_reference: str | None | _UnsetType = _UNSET,
        source_values: Mapping[ResearchSource, str] | _UnsetType = _UNSET,
        research_status: str | None | _UnsetType = _UNSET,
        processed_at: datetime | None = None,
    ) -> None:
        workbook: OpenpyxlWorkbook | None = None
        try:
            workbook = self._load_or_create()
            worksheet = workbook.active
            columns = self._ensure_schema(workbook)
            inquiry_col = columns[INQUIRY_ID_HEADER]
            rows_by_id = self._rows_by_inquiry_id(worksheet, inquiry_col)
            row = rows_by_id.get(inquiry_id, worksheet.max_row + 1)
            worksheet.cell(row=row, column=inquiry_col).value = inquiry_id

            values: dict[str, object | None | _UnsetType] = {
                IMPORTANCE_HEADER: importance_display_value(importance_raw),
                MPN_HEADER: mpn,
                BRAND_HEADER: brand,
                QUANTITY_HEADER: quantity,
                STOCK_HEADER: stock_label,
                ESTIMATED_TOTAL_HEADER: (
                    estimated_total
                    if isinstance(estimated_total, _UnsetType)
                    else _decimal_text(estimated_total)
                ),
                MARKET_REFERENCE_HEADER: market_reference,
                REMARKS_HEADER: remarks,
                PROCESSED_AT_HEADER: _iso_timestamp(processed_at or datetime.now(timezone.utc)),
                RESEARCH_STATUS_HEADER: research_status,
            }
            if not isinstance(source_values, _UnsetType):
                values.update(
                    {
                        header: source_values.get(source, "无结果")
                        for source, header in SOURCE_HEADERS.items()
                    }
                )
            for header, value in values.items():
                if isinstance(value, _UnsetType):
                    continue
                worksheet.cell(row=row, column=columns[header]).value = value
            for header in (MARKET_REFERENCE_HEADER, *SOURCE_HEADERS.values()):
                worksheet.cell(row=row, column=columns[header]).alignment = Alignment(
                    wrap_text=True
                )

            self._save_atomically(workbook)
        except ExcelOutputError:
            raise
        except Exception as exc:
            raise ExcelOutputError("unable to update Research workbook") from exc
        finally:
            if workbook is not None:
                workbook.close()

    def read_history(self) -> tuple[ResearchHistoryRecord, ...]:
        """Read recognized workbook rows without migrating or modifying the file."""
        if not self.path.is_file():
            return ()
        workbook = None
        try:
            workbook = load_workbook(self.path, read_only=True, data_only=True)
            worksheet = workbook.active
            headers = tuple(cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1)))
            if len(set(headers)) != len(headers):
                raise ExcelConsistencyError("duplicate Excel header")
            recognized = {
                frozenset(CANONICAL_HEADERS),
                frozenset(_OLD_CANONICAL_HEADERS),
                frozenset(_PREVIOUS_HEADERS),
                frozenset(_LEGACY_HEADERS),
            }
            if frozenset(headers) not in recognized:
                raise ExcelConsistencyError("unrecognized Research workbook schema")
            columns = {header: index for index, header in enumerate(headers) if header is not None}
            if INQUIRY_ID_HEADER not in columns:
                raise ExcelConsistencyError("Research workbook has no inquiry identity")
            records = []
            inquiry_ids = set()
            for values in worksheet.iter_rows(min_row=2, values_only=True):
                inquiry_id = _cell(values, columns.get(INQUIRY_ID_HEADER))
                if not isinstance(inquiry_id, str) or not inquiry_id.strip():
                    continue
                if inquiry_id in inquiry_ids:
                    raise ExcelConsistencyError(
                        f"duplicate {INQUIRY_ID_HEADER} rows for inquiry_id"
                    )
                inquiry_ids.add(inquiry_id)
                sources = tuple(
                    (header, str(_cell(values, columns.get(header)) or "无结果"))
                    for header in SOURCE_HEADERS.values()
                )
                records.append(ResearchHistoryRecord(
                    inquiry_id=inquiry_id,
                    mpn=_string(_cell(values, columns.get(MPN_HEADER))),
                    brand=_string(_cell(values, columns.get(BRAND_HEADER))),
                    quantity=_cell(values, columns.get(QUANTITY_HEADER)),
                    importance_raw=_string(_cell(values, columns.get(IMPORTANCE_HEADER))),
                    stock_label=_string(_cell(values, columns.get(STOCK_HEADER))),
                    estimated_total=_cell(values, columns.get(ESTIMATED_TOTAL_HEADER, columns.get(_OLD_TOTAL_HEADER))),
                    market_reference=_string(_cell(values, columns.get(MARKET_REFERENCE_HEADER))),
                    source_values=sources,
                    remarks=_string(_cell(values, columns.get(REMARKS_HEADER))),
                    processed_at=_parse_timestamp(_cell(values, columns.get(PROCESSED_AT_HEADER))),
                    research_status=_string(_cell(values, columns.get(RESEARCH_STATUS_HEADER))),
                ))
            ordered = sorted(
                enumerate(records),
                key=lambda pair: (
                    pair[1].processed_at is None,
                    -pair[1].processed_at.timestamp()
                    if pair[1].processed_at
                    else 0,
                    pair[0],
                ),
            )
            return tuple(record for _, record in ordered)
        except ExcelOutputError:
            raise
        except Exception as exc:
            raise ExcelOutputError("unable to read Research history") from exc
        finally:
            if workbook is not None:
                workbook.close()

    def _load_or_create(self) -> OpenpyxlWorkbook:
        if self.path.exists():
            try:
                return load_workbook(self.path)
            except Exception as exc:
                raise ExcelOutputError("unable to load Research workbook") from exc
        workbook = Workbook()
        workbook.active.title = DEFAULT_SHEET_TITLE
        return workbook

    @staticmethod
    def _rows_by_inquiry_id(worksheet, inquiry_col: int) -> dict[object, int]:
        rows: dict[object, int] = {}
        for row in range(2, worksheet.max_row + 1):
            value = worksheet.cell(row=row, column=inquiry_col).value
            if value is None:
                continue
            if value in rows:
                raise ExcelConsistencyError(
                    f"duplicate {INQUIRY_ID_HEADER} rows for inquiry_id"
                )
            rows[value] = row
        return rows

    def _ensure_schema(self, workbook: OpenpyxlWorkbook) -> dict[str, int]:
        worksheet = workbook.active
        blank = (
            worksheet.max_row == 1
            and worksheet.max_column == 1
            and worksheet.cell(row=1, column=1).value is None
        )
        if blank:
            for index, header in enumerate(CANONICAL_HEADERS, start=1):
                worksheet.cell(row=1, column=index, value=header)
        else:
            headers = tuple(
                worksheet.cell(row=1, column=column).value
                for column in range(1, worksheet.max_column + 1)
            )
            if len(set(headers)) != len(headers):
                raise ExcelConsistencyError("duplicate Excel header")
            known_sets = {
                frozenset(CANONICAL_HEADERS),
                frozenset(_OLD_CANONICAL_HEADERS),
                frozenset(_PREVIOUS_HEADERS),
                frozenset(_LEGACY_HEADERS),
            }
            if frozenset(headers) not in known_sets or len(headers) not in {
                len(CANONICAL_HEADERS),
                len(_OLD_CANONICAL_HEADERS),
                len(_PREVIOUS_HEADERS),
                len(_LEGACY_HEADERS),
            }:
                raise ExcelConsistencyError("unrecognized Research workbook schema")
            source_columns = {
                header: index for index, header in enumerate(headers, start=1)
            }
            inquiry_col = source_columns[INQUIRY_ID_HEADER]
            self._rows_by_inquiry_id(worksheet, inquiry_col)
            if headers != CANONICAL_HEADERS:
                self._normalize_schema(worksheet, headers)

        columns = {
            header: index for index, header in enumerate(CANONICAL_HEADERS, start=1)
        }
        for header in VISIBLE_HEADERS:
            worksheet.column_dimensions[
                worksheet.cell(row=1, column=columns[header]).column_letter
            ].hidden = False
        worksheet.column_dimensions[
            worksheet.cell(row=1, column=columns[INQUIRY_ID_HEADER]).column_letter
        ].hidden = True
        worksheet.column_dimensions[
            worksheet.cell(row=1, column=columns[RESEARCH_STATUS_HEADER]).column_letter
        ].hidden = True
        return columns

    @staticmethod
    def _normalize_schema(worksheet, headers: tuple[object, ...]) -> None:
        source_columns = {
            header: index for index, header in enumerate(headers, start=1)
        }
        snapshots: dict[tuple[int, object], tuple[object, object, object, object]] = {}
        for row in range(1, worksheet.max_row + 1):
            for header, column in source_columns.items():
                cell = worksheet.cell(row=row, column=column)
                snapshots[(row, header)] = (
                    cell.value,
                    copy(cell._style),
                    copy(cell.hyperlink),
                    copy(cell.comment),
                )

        aliases = {ESTIMATED_TOTAL_HEADER: _OLD_TOTAL_HEADER}
        original_max_column = worksheet.max_column
        for column, header in enumerate(CANONICAL_HEADERS, start=1):
            source_header = header if header in source_columns else aliases.get(header)
            for row in range(1, worksheet.max_row + 1):
                target = worksheet.cell(row=row, column=column)
                state = (
                    snapshots.get((row, source_header))
                    if source_header is not None
                    else None
                )
                if state is None:
                    target.value = header if row == 1 else None
                    target.hyperlink = None
                    target.comment = None
                    continue
                value, style, hyperlink, comment = state
                target.value = header if row == 1 else value
                target._style = copy(style)
                target.hyperlink = copy(hyperlink)
                target.comment = copy(comment)

        if original_max_column > len(CANONICAL_HEADERS):
            worksheet.delete_cols(
                len(CANONICAL_HEADERS) + 1,
                original_max_column - len(CANONICAL_HEADERS),
            )

    def _save_atomically(self, workbook: OpenpyxlWorkbook) -> None:
        temp_path: Path | None = None
        try:
            self.path.parent.mkdir(parents=True, exist_ok=True)
            fd, temp_name = tempfile.mkstemp(
                prefix=f".{self.path.stem}.",
                suffix=self.path.suffix or ".xlsx",
                dir=self.path.parent,
            )
            os.close(fd)
            temp_path = Path(temp_name)
            workbook.save(temp_path)
            os.replace(temp_path, self.path)
        except Exception as exc:
            if temp_path is not None:
                try:
                    temp_path.unlink(missing_ok=True)
                except OSError:
                    pass
            raise ExcelWriteError("unable to save Research workbook") from exc


def _cell(values, index):
    return values[index] if index is not None and index < len(values) else None


def _string(value):
    return value if isinstance(value, str) else (str(value) if value is not None else None)


def _iso_timestamp(value: datetime) -> str:
    if value.tzinfo is None or value.utcoffset() is None:
        raise ValueError("processed_at must be timezone-aware")
    return value.astimezone(timezone.utc).isoformat(timespec="seconds")


def _parse_timestamp(value) -> datetime | None:
    if isinstance(value, datetime):
        if value.tzinfo is None or value.utcoffset() is None:
            return None
        return value.astimezone(timezone.utc)
    if not isinstance(value, str) or not value.strip():
        return None
    try:
        parsed = datetime.fromisoformat(value)
    except ValueError:
        return None
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        return None
    return parsed.astimezone(timezone.utc)
