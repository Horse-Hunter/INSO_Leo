"""Production GUI adapter over the canonical Workflow and Research runtimes."""

from __future__ import annotations

import json
import logging
import os
import subprocess
import sys
import threading
import uuid
from dataclasses import replace
from datetime import timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from src.gui.contracts import (
    DiagnosticSnapshot,
    GuiBackend,
    HealthItem,
    HealthReport,
    LogEntry,
    Order,
    OrderStatus,
    RunSession,
    RunState,
    SourceDetail,
)
from src.gui.state import utc_now
from src.research import ResearchInput, ResearchResult
from src.research.runtime import (
    ResearchRuntimeConfigError,
    assess_readiness,
    build_production_research_service,
    load_runtime_config,
)
from src.sheets import WorksheetIdentity
from src.sheets.google_oauth import build_read_only_google_sheets_service
from src.sheets.google_reader import GoogleSheetsRowReader
from src.workflow import (
    WorkflowPoller,
    WorkflowRuntime,
    WorkflowStateStore,
    WorkflowStatus,
    WorkflowWorker,
)

log = logging.getLogger(__name__)


class _Observer:
    def __init__(self, service, seen):
        self.service, self.seen = service, seen

    def execute(self, item: ResearchInput) -> ResearchResult:
        self.seen(item.inquiry_id)
        return self.service.execute(item)


class ProductionBackend(GuiBackend):
    """Own one GUI run session and stop it only at safe Workflow boundaries."""

    def __init__(
        self, config_path=None, production_config_path=None, *, cdp_probe=None
    ):
        self.root = Path(__file__).resolve().parents[2]
        self.config_path = Path(config_path or self.root / "runtime/research.json")
        self.production_path = Path(
            production_config_path or self.root / "runtime/production.json"
        )
        self.cdp_probe = cdp_probe
        self._lock = threading.RLock()
        self._stop = threading.Event()
        self._thread = None
        self._closed = False
        self._run_id = None
        self._state = RunState.STOPPED
        self._started = self._stopped = self._last_poll = None
        self._inquiries = []
        self._results = ()
        self._store = None
        self._excel = None
        self._manual_inquiries = set()
        try:
            configured_excel = load_runtime_config(self.config_path).excel_output_path
            self._excel = (
                configured_excel
                if configured_excel.is_absolute()
                else self.root / configured_excel
            )
        except ResearchRuntimeConfigError as exc:
            log.debug("Research output path unavailable (%s)", type(exc).__name__)
        self._logs = []
        self._status_callbacks = []
        self._log_callbacks = []
        self._health = HealthReport(
            tuple(HealthItem(n, "未知", "尚未检查") for n in self._components()),
            "需要人工处理",
        )

    @staticmethod
    def _components():
        return ("Google Sheets", "Browser/CDP", "Credential", "Research")

    @property
    def run_id(self):
        with self._lock:
            return self._run_id

    def start(self):
        with self._lock:
            if self._closed or self._state in (
                RunState.RUNNING,
                RunState.STOPPING_AFTER_CYCLE,
            ):
                return
            self._run_id = "run_" + str(uuid.uuid4())
            self._started = utc_now()
            self._stopped = None
            self._inquiries = []
            self._manual_inquiries.clear()
            self._results = ()
            self._last_poll = None
            self._stop.clear()
            self._state = RunState.RUNNING
            self._thread = threading.Thread(
                target=self._run, name="production-launcher", daemon=False
            )
            self._thread.start()
            self._notify()

    def request_stop_after_cycle(self):
        with self._lock:
            if self._state is RunState.RUNNING:
                self._state = RunState.STOPPING_AFTER_CYCLE
                self._stop.set()
                self._notify()

    def _run(self):
        try:
            cfg = json.loads(self.production_path.read_text(encoding="utf-8"))
            if (
                not cfg.get("spreadsheet_id")
                or not cfg.get("worksheet_titles")
                or not cfg.get("client_secret_file")
            ):
                raise ValueError(
                    "production.json requires spreadsheet_id, worksheet_titles and client_secret_file"
                )
            rc = load_runtime_config(self.config_path)
            if not rc.excel_output_path.is_absolute():
                rc = replace(rc, excel_output_path=self.root / rc.excel_output_path)
            readiness = assess_readiness(rc.cdp.cdp_url, cdp_probe=self.cdp_probe)
            if not readiness.ready:
                raise RuntimeError("Research readiness requires manual handling")
            client = Path(cfg["client_secret_file"])
            client = client if client.is_absolute() else self.root / client
            reader = GoogleSheetsRowReader(
                build_read_only_google_sheets_service(client)
            )
            worksheets = tuple(
                WorksheetIdentity(cfg["spreadsheet_id"], title)
                for title in cfg["worksheet_titles"]
            )
            db = Path(cfg.get("sqlite_path", "runtime/production/workflow.sqlite3"))
            db = db if db.is_absolute() else self.root / db
            self._excel = (
                rc.excel_output_path
                if rc.excel_output_path.is_absolute()
                else self.root / rc.excel_output_path
            )
            self._store = WorkflowStateStore(db)
            research = build_production_research_service(rc, cdp_probe=self.cdp_probe)
            worker = WorkflowWorker(
                self._store,
                _Observer(research, self._seen, self._manual_review),
                brand_updater=None,
            )
            runtime = WorkflowRuntime(
                WorkflowPoller(self._store, reader), worker, worksheets
            )
            self._set_health(("正常", "已连接", "正常", "正常"), "正常")

            def poll_loop():
                try:
                    while not self._stop.is_set():
                        if self._stop.is_set():
                            return
                        self._last_poll = utc_now()
                        runtime.run_poll(now=self._last_poll)
                        self._set_health(("正常", "已连接", "正常", "正常"), "正常")
                        self._stop.wait(runtime.poll_interval.total_seconds())
                except Exception as exc:  # noqa: BLE001 - fail closed at runtime boundary
                    self._runtime_error(exc)

            def worker_loop():
                try:
                    while not self._stop.is_set():
                        runtime.worker.process_due_one()
                        self._refresh()
                        self._stop.wait(runtime.worker_idle_interval.total_seconds())
                except Exception as exc:  # noqa: BLE001 - fail closed at runtime boundary
                    self._runtime_error(exc)

            poll_thread = threading.Thread(
                target=poll_loop, name="production-sheets-poller"
            )
            worker_thread = threading.Thread(
                target=worker_loop, name="production-research-worker"
            )
            poll_thread.start()
            worker_thread.start()
            while poll_thread.is_alive() or worker_thread.is_alive():
                self._refresh()
                poll_thread.join(timeout=0.1)
                worker_thread.join(timeout=0.1)
        except Exception as exc:  # noqa: BLE001 - fail closed at runtime boundary
            log.warning("Production runtime stopped (%s)", type(exc).__name__)
            self._append_log(
                "ERROR",
                f"需要人工处理：{type(exc).__name__}；请检查本地配置、授权、CDP 与人工验证状态",
            )
            self._set_health(("异常", "异常", "异常", "需要人工处理"), "需要人工处理")
            with self._lock:
                self._state = RunState.MANUAL_REVIEW
        finally:
            self._refresh()
            with self._lock:
                if self._state is not RunState.MANUAL_REVIEW:
                    self._state = RunState.STOPPED
                self._stopped = utc_now()
                self._notify()

    def _seen(self, inquiry):
        with self._lock:
            if inquiry not in self._inquiries:
                self._inquiries.append(inquiry)

    def _refresh(self):
        if not self._store:
            return
        from openpyxl import load_workbook

        items = {i.inquiry_id: i for i in self._store.all_items()}
        ids = [i for i in self._inquiries if i in items]
        rows = {}
        if self._excel and self._excel.is_file() and ids:
            wb = load_workbook(self._excel, read_only=True, data_only=True)
            try:
                ws = wb.active
                cols = {ws.cell(1, c).value: c - 1 for c in range(1, ws.max_column + 1)}
                for vals in ws.iter_rows(min_row=2, values_only=True):
                    key = vals[cols["_inquiry_id"]] if "_inquiry_id" in cols else None
                    if key in ids:
                        rows[key] = {h: vals[c] for h, c in cols.items()}
            finally:
                wb.close()
        source_names = ("INSO", "Findchips", "华强", "立创", "正能量")
        results = []
        for key in ids:
            item = items[key]
            row = rows.get(key, {})
            status = {
                WorkflowStatus.COMPLETED: (
                    OrderStatus.PARTIAL
                    if item.research_status == "PARTIAL_SUCCESS"
                    else OrderStatus.COMPLETED
                ),
                WorkflowStatus.MANUAL_REVIEW: OrderStatus.MANUAL_REVIEW,
                WorkflowStatus.FAILED: OrderStatus.MANUAL_REVIEW,
            }.get(item.status, OrderStatus.PENDING)
            if key in self._manual_inquiries:
                status = OrderStatus.MANUAL_REVIEW
            results.append(
                Order(
                    key,
                    str(row.get("型号") or item.mpn),
                    row.get("品牌") or item.brand,
                    _integer(row.get("数量"), item.quantity),
                    str(row.get("货量标识") or "待验证"),
                    _decimal(row.get("市场最低参考价")),
                    _decimal(row.get("预估订单总价")),
                    status,
                    tuple(
                        SourceDetail(n, str(row.get(n) or "无结果"))
                        for n in source_names
                    ),
                    str(row.get("备注") or item.last_error or ""),
                    self._run_id or "",
                )
            )
        with self._lock:
            self._results = tuple(results)

    def get_status(self):
        with self._lock:
            items = self._store.all_items() if self._store else ()
            return RunSession(
                self._run_id,
                self._state,
                self._started,
                self._stopped,
                len(self._results),
                sum(
                    i.status is WorkflowStatus.COMPLETED
                    and i.inquiry_id in self._inquiries
                    for i in items
                ),
                sum(i.status is WorkflowStatus.RESEARCHING for i in items),
                sum(
                    i.status in (WorkflowStatus.QUEUED, WorkflowStatus.RETRY_WAIT)
                    for i in items
                ),
                self._last_poll + timedelta(minutes=15) if self._last_poll else None,
            )

    def get_current_run_results(self):
        with self._lock:
            return self._results

    def get_health(self):
        return self._health

    def get_logs(self):
        with self._lock:
            return tuple(self._logs)

    def get_diagnostics(self):
        with self._lock:
            return DiagnosticSnapshot(
                self._run_id,
                max(0, (utc_now() - self._started).total_seconds())
                if self._started
                else 0,
                0,
                "running" if self._thread and self._thread.is_alive() else "stopped",
                self._last_poll,
            )

    def open_excel(self):
        if self._excel and self._excel.exists():
            os.startfile(self._excel)

    def open_results_dir(self):
        path = self._excel.parent if self._excel else self.root / "runtime"
        if sys.platform == "win32":
            os.startfile(path)
        else:
            subprocess.Popen(["xdg-open", str(path)])

    def on_status_change(self, callback):
        with self._lock:
            if callback is None:
                self._status_callbacks.clear()
            elif callback not in self._status_callbacks:
                self._status_callbacks.append(callback)

    def on_log(self, callback):
        with self._lock:
            if callback is None:
                self._log_callbacks.clear()
            elif callback not in self._log_callbacks:
                self._log_callbacks.append(callback)

    def shutdown(self):
        with self._lock:
            if self._closed:
                return
            self._closed = True
            self._stop.set()
            thread = self._thread
        if thread and thread is not threading.current_thread():
            thread.join()
        with self._lock:
            self._status_callbacks.clear()
            self._log_callbacks.clear()

    def _notify(self):
        status = self.get_status()
        for cb in tuple(self._status_callbacks):
            cb(status)

    def _append_log(self, level, message):
        entry = LogEntry(utc_now(), level, message)
        with self._lock:
            self._logs = (self._logs + [entry])[-1000:]
            callbacks = tuple(self._log_callbacks)
        for cb in callbacks:
            cb(entry)

    def _set_health(self, statuses, overall):
        self._health = HealthReport(
            tuple(
                HealthItem(n, s)
                for n, s in zip(self._components(), statuses, strict=True)
            ),
            overall,
        )


def _integer(value, fallback):
    try:
        return int(value)
    except (ValueError, TypeError):
        return fallback if isinstance(fallback, int) else 0


def _decimal(value):
    if value is None:
        return None
    text = str(value).strip().splitlines()[0].replace("¥", "").replace(",", "")
    try:
        return Decimal(text) if text else None
    except InvalidOperation:
        return None
