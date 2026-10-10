"""PI No. from Excel only, using the approved read-only receiver."""
import re
from dataclasses import dataclass, field
from email.utils import parsedate_to_datetime
from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook

from .inspection import MAX_BYTES, MAX_CELLS, inspect_order_mail

SAMPLE_DATES = frozenset({"2026-10-09T16:45:52+08:00", "2026-10-07T12:26:16+08:00"})
_PI_LABEL = re.compile(r"\bPI\s*(?:NO\.?|NUMBER)\s*[:：]?", re.IGNORECASE)
_PI_VALUE = re.compile(r"SHAWN\d{8}-\d{2,}\b")


class PiError(ValueError):
    """Closed reason codes only; never include raw values."""


@dataclass(frozen=True)
class PiOrder:
    sample_number: int
    pi_no: str = field(repr=False)


def extract_pi(payload: bytes) -> str:
    with ZipFile(BytesIO(payload)) as archive:
        if sum(i.file_size for i in archive.infolist()) > MAX_BYTES * 4:
            raise PiError("EXCEL_TOO_LARGE")
    wb = load_workbook(BytesIO(payload), data_only=False, keep_links=False)
    try:
        locations = []
        for sheet in wb.worksheets:
            if sheet.max_row * sheet.max_column > MAX_CELLS:
                raise PiError("EXCEL_STRUCTURE_UNCONFIRMED")
            for row in sheet.iter_rows():
                for cell in row:
                    text = cell.value if isinstance(cell.value, str) else ""
                    label = _PI_LABEL.search(text)
                    if not label:
                        continue
                    tail = text[label.end():].strip()
                    if not tail:
                        adjacent = sheet.cell(cell.row, cell.column + 1)
                        if adjacent.data_type == "f" or not isinstance(adjacent.value, str):
                            raise PiError("PI_STRUCTURE_UNCONFIRMED")
                        tail = adjacent.value
                    candidates = _PI_VALUE.findall(tail)
                    if len(candidates) != 1 or cell.data_type == "f":
                        raise PiError("PI_MISSING_OR_AMBIGUOUS")
                    locations.append(candidates[0])
        if len(locations) != 1:
            raise PiError("PI_MISSING_OR_AMBIGUOUS")
        return locations[0]
    finally:
        wb.close()


def read_sample_pi_orders(*, inspect=inspect_order_mail):
    """Only the two Phase1-confirmed dated messages, never other new emails."""
    values = []
    def collect(message):
        try:
            date = parsedate_to_datetime(str(message.get("Date", ""))).isoformat()
        except (ValueError, TypeError):
            return
        if date not in SAMPLE_DATES:
            return
        excels = [part for part in message.walk()
                  if str(part.get_filename() or "").lower().endswith(".xlsx")]
        if len(excels) != 1:
            raise PiError("EXCEL_ATTACHMENT_AMBIGUOUS")
        values.append((date, extract_pi(excels[0].get_payload(decode=True) or b"")))
    errors = []
    def collect_safely(message):
        try:
            collect(message)
        except PiError as exc:
            errors.append(str(exc))
        except Exception:  # noqa: BLE001 - malformed attachment never leaks content
            errors.append("EXCEL_STRUCTURE_UNCONFIRMED")
    report = inspect(_message_consumer=collect_safely)
    if report.status != "SUCCESS" or report.flags_unchanged is not True:
        raise PiError("MAIL_READ_NOT_CONFIRMED")
    if errors:
        raise PiError(errors[0])
    if len(values) != 2 or {date for date, _ in values} != SAMPLE_DATES:
        raise PiError("CONFIRMED_SAMPLES_UNAVAILABLE")
    values.sort(reverse=True)
    return tuple(PiOrder(i, value) for i, (_, value) in enumerate(values, 1))
