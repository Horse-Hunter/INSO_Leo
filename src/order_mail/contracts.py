"""Strict in-memory contract parsing for one selected approved order."""
import re
from dataclasses import dataclass, field
from datetime import date, timedelta
from decimal import Decimal
from email.utils import parsedate_to_datetime
from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook

from .inspection import MAX_BYTES, MAX_CELLS, inspect_order_mail
from .pi_orders import SAMPLE_DATES, extract_pi


class OrderError(ValueError):
    def __init__(self, code, *, field=None, pi_no=None, part_number=None):
        super().__init__(code)
        self.code, self.field = code, field
        self.pi_no, self.part_number = pi_no, part_number


@dataclass(frozen=True, repr=False)
class ContractLine:
    part_number: str
    lead_time: str
    brand: str
    dc: str
    quantity: int
    unit_price: Decimal
    delivery_date: str


@dataclass(frozen=True)
class ContractOrder:
    pi_no: str = field(repr=False)
    lines: tuple[ContractLine, ...] = field(repr=False)
    today: date


def delivery_date(value: str, today: date) -> str:
    text = value.strip().upper()
    if text == "1-3 DAYS":
        days = 7
    else:
        match = re.fullmatch(r"(\d+)-(\d+) WEEKS", text)
        if not match or not 0 < int(match[1]) <= int(match[2]):
            raise OrderError("LEAD_TIME", field="L/T")
        days = int(match[2]) * 7 + 7
    return (today + timedelta(days=days)).isoformat()


def _header(value):
    return re.sub(r"\s+", "", str(value or '').upper())


_HEADERS = {"PARTNUMBER": "part_number", "L/T": "lead_time", "BRAND": "brand",
            "DC": "dc", "QTY(PCS)": "quantity", "QTY": "quantity",
            "UNITPRICE(￥)": "unit_price", "UNITPRICE(¥)": "unit_price"}
_LABELS = {"part_number": "PART NUMBER", "lead_time": "L/T", "brand": "BRAND",
           "dc": "DC", "quantity": "QTY", "unit_price": "Unit Price(￥)"}


def parse_contract(payload: bytes, *, today: date) -> ContractOrder:
    pi = None
    try:
        pi = extract_pi(payload)
        with ZipFile(BytesIO(payload)) as archive:
            if sum(x.file_size for x in archive.infolist()) > MAX_BYTES * 4:
                raise OrderError("STRUCTURE")
        wb = load_workbook(BytesIO(payload), data_only=False, keep_links=False)
        try:
            headers = []
            for sheet in wb.worksheets:
                if sheet.max_row * sheet.max_column > MAX_CELLS:
                    raise OrderError("STRUCTURE")
                for row in sheet:
                    fields = {}
                    duplicate = False
                    for cell in row:
                        key = _HEADERS.get(_header(cell.value))
                        if key:
                            if key in fields: duplicate = True
                            fields[key] = cell.column
                    if "part_number" in fields:
                        if duplicate or set(fields) != set(_LABELS):
                            missing = next((k for k in _LABELS if k not in fields), None)
                            raise OrderError("REQUIRED", field=_LABELS.get(missing, "明细表头"))
                        headers.append((sheet, row[0].row, fields))
            if len(headers) != 1:
                raise OrderError("STRUCTURE", field="明细表头")
            sheet, header_row, columns = headers[0]
            lines, ended = [], False
            for r in range(header_row + 1, sheet.max_row + 1):
                first = sheet.cell(r, columns["part_number"])
                # Confirmed template ends with merged TOTAL PAYMENT footer.
                if isinstance(first.value, str) and "TOTAL" in first.value.upper() and any(
                        m.min_row == r == m.max_row and m.min_col == columns["part_number"]
                        and m.max_col >= max(columns.values()) for m in sheet.merged_cells.ranges):
                    ended = True
                    break
                values = {}
                for key, col in columns.items():
                    cell = sheet.cell(r, col)
                    if cell.data_type == "f" or cell.value is None or cell.value == "":
                        raise OrderError("REQUIRED", field=_LABELS[key], part_number=values.get("part_number"))
                    if key not in {"quantity", "unit_price"}:
                        if not isinstance(cell.value, str) or not cell.value.strip() or any(
                                ch in cell.value for ch in "\r\n"):
                            raise OrderError("TYPE", field=_LABELS[key], part_number=values.get("part_number"))
                        values[key] = cell.value.strip()
                    else:
                        if isinstance(cell.value, bool) or not isinstance(cell.value, (int, float, Decimal)):
                            raise OrderError("TYPE", field=_LABELS[key], part_number=values.get("part_number"))
                        number = Decimal(str(cell.value))
                        if not number.is_finite() or number <= 0 or (key == "quantity" and number != number.to_integral_value()):
                            raise OrderError("TYPE", field=_LABELS[key], part_number=values.get("part_number"))
                        values[key] = int(number) if key == "quantity" else number
                try:
                    due = delivery_date(values["lead_time"], today)
                except OrderError as exc:
                    exc.part_number = values["part_number"]
                    raise
                lines.append(ContractLine(**values, delivery_date=due))
            if not ended or not lines:
                raise OrderError("STRUCTURE", field="明细表格")
            return ContractOrder(pi, tuple(lines), today)
        finally:
            wb.close()
    except OrderError as exc:
        exc.pi_no = pi
        raise
    except Exception:  # noqa: BLE001 - raw workbook errors never exposed
        raise OrderError("STRUCTURE", pi_no=pi) from None


def read_selected_contract(sample_number: int, *, today: date, inspect=inspect_order_mail):
    dates = sorted(SAMPLE_DATES, reverse=True)
    if sample_number not in (1, 2):
        raise OrderError("SELECTION")
    selected = dates[sample_number - 1]
    payloads = []
    def collect(message):
        if parsedate_to_datetime(str(message.get("Date", ""))).isoformat() != selected:
            return
        parts = [p for p in message.walk() if str(p.get_filename() or '').lower().endswith('.xlsx')]
        payloads.append(parts[0].get_payload(decode=True) if len(parts) == 1 else None)
    report = inspect(selected_date=selected, _message_consumer=collect)
    if report.status != "SUCCESS" or report.flags_unchanged is not True:
        raise OrderError("MAIL")
    if len(payloads) != 1 or payloads[0] is None:
        raise OrderError("STRUCTURE")
    return parse_contract(payloads[0], today=today)
