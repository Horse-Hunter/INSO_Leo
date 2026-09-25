from pathlib import Path

from openpyxl import load_workbook

from src.research.contracts import ResearchStatus
from src.research.excel_output import (
    IMPORTANCE_HEADER,
    ExcelWriteError,
    ResearchExcelOutput,
)
from src.research.finalize import finalize_research_result


class FailingOutput:
    def upsert(
        self,
        inquiry_id: str,
        *,
        importance_raw: str | None,
        remarks: str | None = None,
        research_status: str | None = None,
    ) -> None:
        raise ExcelWriteError("synthetic failure")


def test_legacy_finalizer_is_not_exported_from_package_api() -> None:
    from src import research

    assert not hasattr(research, "finalize_research_result")
    assert "finalize_research_result" not in research.__all__


def test_success_requires_successful_persistence(tmp_path: Path) -> None:
    output = ResearchExcelOutput(tmp_path / "调研价格.xlsx")

    result = finalize_research_result(
        output=output,
        inquiry_id="inq_1",
        importance_raw="A",
        status=ResearchStatus.SUCCESS,
    )

    assert result.status is ResearchStatus.SUCCESS
    assert (tmp_path / "调研价格.xlsx").exists()


def test_excel_failure_becomes_retryable_failure() -> None:
    result = finalize_research_result(
        output=FailingOutput(),  # type: ignore[arg-type]
        inquiry_id="inq_1",
        importance_raw="A",
        status=ResearchStatus.SUCCESS,
    )

    assert result.status is ResearchStatus.RETRYABLE_FAILURE
    assert result.inquiry_id == "inq_1"


def test_exception_reason_and_importance_are_persisted(tmp_path: Path) -> None:
    path = tmp_path / "调研价格.xlsx"
    result = finalize_research_result(
        output=ResearchExcelOutput(path),
        inquiry_id="inq_1",
        importance_raw="C",
        status=ResearchStatus.EXCEPTION,
        remarks="所有价格来源均无报价",
    )

    worksheet = load_workbook(path).active
    headers = {
        worksheet.cell(1, column).value: column
        for column in range(1, worksheet.max_column + 1)
    }

    assert result.status is ResearchStatus.EXCEPTION
    assert worksheet.cell(2, headers[IMPORTANCE_HEADER]).value == "C"
    assert worksheet.cell(2, headers["备注"]).value == "所有价格来源均无报价"
