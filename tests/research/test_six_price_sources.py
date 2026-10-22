"""Owner six-source rules, using synthetic captures only (no network/runtime)."""

import json
from dataclasses import replace
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font

from src.research.aggregation import aggregate_price_results
from src.research.bom_ai import (
    BomAiAdapter,
    BomAiCapture,
    BomAiPriceRecord,
    parse_bom_ai_price_records,
)
from src.research.contracts import ResearchInput, ResearchStatus
from src.research.excel_output import CANONICAL_HEADERS, ResearchExcelOutput
from src.research.findchips import (
    FindchipsAdapter,
    FindchipsPage,
    parse_findchips_offers,
)
from src.research.fx import UsdRmbQuote
from src.research.hqew import HqewAdapter, HqewPage
from src.research.icnet import (
    IcNetAdapter,
    IcNetPageUnavailable,
    parse_icnet_reference_price,
)
from src.research.inso_history import (
    InsoHistoryAdapter,
    InsoHistoryCapture,
    InsoHistoryRecord,
)
from src.research.lcsc import LcscAdapter, LcscPage, LcscPriceTier, LcscProduct
from src.research.service import ResearchService
from src.research.source_contracts import (
    PRICE_SOURCES,
    ResearchSource,
    SourceOutcome,
    calendar_month_cutoff,
    format_source_result,
)
from tests.research.test_aggregation_service import FakePrice, _complete, _result
from tests.research.test_findchips import _page, _row
from tests.research.test_hqew import _row as hqew_row
from tests.research.test_icnet import FakeClient

NOW = datetime(2026, 10, 22, 12, tzinfo=UTC)
IC_HTML = (Path(__file__).parent / 'fixtures/icnet_results_minimal.html').read_text(encoding='utf-8')
HISTORICAL = ('findchips', 'hqew', 'bom.ai', 'inso')


class Fx:
    def get_quote(self):
        return UsdRmbQuote(Decimal(7), NOW, 'synthetic-test-only')

    def get_hkd_rmb_rate(self):
        return Decimal('0.9')


def history(source, records, *, stock=True):
    """Exercise actual parsers/adapters, including quotation dates and FX."""
    if source == 'findchips':
        html = _page(*(_row('ABC', '1' if stock else '0',
            json.dumps([[1, 'USD', str(price)]]),
            when.isoformat() if when else None) for price, when in records))
        return FindchipsAdapter(SimpleNamespace(fetch_first_page=lambda _: FindchipsPage(html, 'https://www.findchips.com/search/ABC', NOW)), Fx(), clock=lambda: NOW).search('ABC', 99)
    if source == 'hqew':
        html = ''.join(hqew_row('ABC', str(price), when.isoformat() if when else None) for price, when in records)
        return HqewAdapter(SimpleNamespace(fetch_first_page=lambda _: HqewPage(html, 'https://p.hqew.com/yunquote/ABC.html', NOW)), clock=lambda: NOW).search('ABC', 99)
    if source == 'bom.ai':
        capture = BomAiCapture(tuple(BomAiPriceRecord('ABC', Decimal(price), when) for price, when in records), 'https://www.bom.ai/search/ABC', NOW)
        return BomAiAdapter(SimpleNamespace(fetch_price_records=lambda _: capture), clock=lambda: NOW).search('ABC', 99)
    capture = InsoHistoryCapture(tuple(InsoHistoryRecord(Decimal(price), 'RMB', when, 'ABC') for price, when in records), 'https://synthetic.invalid/history', NOW)
    return InsoHistoryAdapter(SimpleNamespace(fetch_history=lambda _: capture), Fx(), clock=lambda: NOW).search('ABC', 99)


@pytest.mark.parametrize('source', HISTORICAL)
@pytest.mark.parametrize('months', range(1, 7))
def test_first_valid_calendar_window_one_through_six(source, months):
    result = history(source, [('8', calendar_month_cutoff(NOW, months)),
                              ('1', calendar_month_cutoff(NOW, months) - timedelta(days=1))])
    assert result.outcome is SourceOutcome.SUCCESS
    assert result.price_candidate.age_months == months
    assert result.price_candidate.raw_price == Decimal(8)
    if months > 1:
        assert ('两个月', '三个月', '四个月', '五个月', '六个月')[months - 2] in format_source_result(result)


@pytest.mark.parametrize('source', HISTORICAL)
def test_recent_quote_wins_over_cheaper_older_quote(source):
    result = history(source, [('10', NOW - timedelta(days=2)), ('1', calendar_month_cutoff(NOW, 2))])
    assert result.price_candidate.raw_price == Decimal(10)
    assert result.price_candidate.age_months == 1


@pytest.mark.parametrize('source', HISTORICAL)
@pytest.mark.parametrize('when', [calendar_month_cutoff(NOW, 6) - timedelta(seconds=1), NOW + timedelta(seconds=1)])
def test_future_or_beyond_six_months_is_not_a_valid_price(source, when):
    result = history(source, [('1', when)])
    assert result.outcome is SourceOutcome.NO_VALID_PRICE
    assert result.price_candidate is None


def reference(text):
    return '<div class="navigation"><div class="cankaoPrice"><a><span>' + text + '</span></a></div></div>'


@pytest.mark.parametrize('text,expected', [('1.23', '1.23'), ('￥0.015', '0.015'), ('RMB 2.5元', '2.5'),
    ('参考价：¥1,234.50/个', '1234.50'), ('', '0'), ('0', '0'), ('￥0.00', '0'), ('暂无', '0'),
    ('NaN', '0'), ('Infinity', '0'), ('-1', '0'), ('1-2元', '0'), ('USD 2', '0'), ('1,2', '0')])
def test_icnet_reference_price_and_brand_stock_are_independent(text, expected):
    baseline = IcNetAdapter(FakeClient(IC_HTML)).search('ABC-123', None, 10)
    client = FakeClient(IC_HTML.replace('<body>', '<body>' + reference(text)))
    actual = IcNetAdapter(client).search('ABC-123', None, 10)
    assert client.queries == ['ABC-123']
    assert actual.resolved_brand == baseline.resolved_brand
    assert actual.stock_label == baseline.stock_label
    fields = {field.key: field.value for field in actual.source_result.evidence.fields}
    old_fields = {field.key: field.value for field in baseline.source_result.evidence.fields}
    assert {k: v for k, v in fields.items() if k != 'reference_price_rmb'} == {k: v for k, v in old_fields.items() if k != 'reference_price_rmb'}
    assert fields['reference_price_rmb'] == Decimal(expected)
    candidate = actual.source_result.price_candidate
    if Decimal(expected):
        assert candidate.normalized_rmb_price == Decimal(expected)
        assert candidate.raw_currency == 'RMB'
    else:
        assert candidate is None
        assert format_source_result(actual.source_result) == '0'
    assert actual.source_result.outcome is SourceOutcome.SUCCESS


@pytest.mark.parametrize('code', ['BROWSER_TIMEOUT', 'LOGIN_REQUIRED', 'MANUAL_VERIFICATION_REQUIRED'])
def test_icnet_technical_failure_is_never_zero(code):
    def fail(_):
        raise IcNetPageUnavailable(code)
    actual = IcNetAdapter(SimpleNamespace(fetch_first_page=fail), clock=lambda: NOW).search('ABC', None, 10)
    assert actual.source_result.outcome is SourceOutcome.SOURCE_UNAVAILABLE
    assert actual.stock_label == '待验证'
    assert format_source_result(actual.source_result) == '暂不可用'
    results = _complete({ResearchSource.IC_NET: actual.source_result})
    assert aggregate_price_results(results, 10).status is ResearchStatus.RETRYABLE_FAILURE


def test_icnet_empty_broken_or_ambiguous_price_does_not_mask_result_failure():
    for html in ('<body>' + reference('3') + '</body>', '<body>登录</body>'):
        result = IcNetAdapter(FakeClient(html)).search('ABC', None, 1)
        assert result.source_result.outcome is SourceOutcome.SOURCE_UNAVAILABLE
    for html in ('<body>' + reference('2') + reference('3') + '</body>',
                 '<body><section>' + reference('2') + '</section></body>',
                 '<body>' + reference('2').replace('<span>', '<span style="display:none">') + '</body>'):
        assert parse_icnet_reference_price(html) == 0


@pytest.mark.parametrize('winner', PRICE_SOURCES)
def test_each_of_six_sources_can_win_and_quantity_is_unchanged(winner):
    results = _complete({source: _result(source, SourceOutcome.SUCCESS, '2' if source is winner else '10') for source in PRICE_SOURCES})
    actual = aggregate_price_results(results, 13)
    assert actual.lowest.source is winner
    assert actual.estimated_total == Decimal(26)
    assert actual.status is ResearchStatus.SUCCESS


@pytest.mark.parametrize('stock', [True, False])
@pytest.mark.parametrize('when', [None, NOW - timedelta(days=400), NOW + timedelta(days=100)])
def test_lcsc_ignores_all_date_windows_but_keeps_stock_and_tiers(stock, when):
    product = LcscProduct('ABC', False, 5 if stock else 0,
        (LcscPriceTier(1, Decimal(9), 'CNY'), LcscPriceTier(10000, Decimal(2), 'CNY')), when)
    page = LcscPage('', 'https://www.lcsc.com/product-detail/C1.html', NOW, product)
    actual = LcscAdapter(SimpleNamespace(fetch_product_page=lambda _: page), Fx(), clock=lambda: NOW).search('ABC', 1)
    assert actual.outcome is SourceOutcome.SUCCESS
    selected = actual.price_candidate if stock else actual.out_of_stock_candidate
    assert selected.normalized_rmb_price == Decimal(2)
    assert (actual.out_of_stock_candidate is None) is stock


def test_findchips_window_and_fallback_keep_existing_global_pool_rules():
    result = history('findchips', [('2', calendar_month_cutoff(NOW, 5))], stock=False)
    assert result.price_candidate is None
    assert result.out_of_stock_candidate.age_months == 5
    assert '五个月' in format_source_result(result)
    pool = _complete({ResearchSource.FINDCHIPS: result})
    assert aggregate_price_results(pool, 2).used_out_of_stock_fallback
    icnet = _result(ResearchSource.IC_NET, SourceOutcome.SUCCESS, '99')
    assert aggregate_price_results(_complete({ResearchSource.FINDCHIPS: result, ResearchSource.IC_NET: icnet}), 2).lowest.source is ResearchSource.IC_NET


def test_undated_current_prices_and_generic_update_dates():
    for source in ('findchips', 'hqew'):
        result = history(source, [('10', None), ('1', calendar_month_cutoff(NOW, 2))])
        assert result.price_candidate.raw_price == Decimal(10)
        assert result.price_candidate.age_months == 1
    html = _page(_row('ABC', '1', '[[1,"USD","1"]]').replace('<tr ', '<tr data-date="2020-01-01" '))
    assert parse_findchips_offers(html, 'ABC')[0].observed_at is None
    dated = '<h3>ABC</h3><data><quotePrice>￥3</quotePrice><quoteDate>2026/04/22 20:00:00</quoteDate></data>'
    assert parse_bom_ai_price_records(dated, 'ABC', captured_at=NOW)[0].observed_at < NOW
    undated = '<h3>ABC</h3><data><quotePrice>￥3</quotePrice></data>'
    assert parse_bom_ai_price_records(undated, 'ABC', captured_at=NOW) == ()
    current = undated.replace('<data>', '<data>当前报价')
    assert parse_bom_ai_price_records(current, 'ABC', captured_at=NOW)[0].observed_at == NOW


def test_service_reuses_icnet_capture_in_six_source_excel(tmp_path):
    client = FakeClient(IC_HTML.replace('<body>', '<body>' + reference('2')))
    by_source = {item.source: item for item in _complete()}
    service = ResearchService(icnet=IcNetAdapter(client),
        findchips=FakePrice(by_source[ResearchSource.FINDCHIPS]), hqew=FakePrice(by_source[ResearchSource.HQEW]),
        lcsc=FakePrice(by_source[ResearchSource.LCSC]), bom_ai=FakePrice(by_source[ResearchSource.BOM_AI]),
        inso=FakePrice(by_source[ResearchSource.INSO]), output=ResearchExcelOutput(tmp_path / 'prices.xlsx'))
    result = service.execute_detailed(ResearchInput('synthetic-six', 'ABC-123', None, 13, 'A'))
    assert client.queries == ['ABC-123']
    assert len(result.price_sources) == 6
    assert result.aggregation.lowest.source is ResearchSource.IC_NET
    assert result.aggregation.estimated_total == Decimal(26)
    assert result.result.status is ResearchStatus.SUCCESS


@pytest.mark.parametrize('reverse', [False, True])
def test_old_five_source_workbook_upgrade_preserves_every_historical_cell(tmp_path, reverse):
    path = tmp_path / 'prices.xlsx'
    old_headers = tuple(header for header in CANONICAL_HEADERS if header != 'IC.net')
    if reverse:
        old_headers = old_headers[::-1]
    wb = Workbook()
    ws = wb.active
    ws.append(old_headers)
    values = {header: 'historical-' + header for header in old_headers}
    values.update(_inquiry_id='synthetic-old', _research_status='PARTIAL_SUCCESS', 数量=10,
                  处理时间=NOW.isoformat(), 预估订单总价='=C2*G2')
    ws.append([values[header] for header in old_headers])
    old_price = ws.cell(2, old_headers.index('INSO') + 1)
    old_price.font = Font(bold=True)
    old_price.comment = Comment('synthetic-note', 'test')
    wb.create_sheet('unrelated')['A1'] = 'preserve-me'
    wb.save(path)
    wb.close()
    output = ResearchExcelOutput(path)
    original_bytes = path.read_bytes()
    assert output.read_history()[0].inquiry_id == 'synthetic-old'
    assert path.read_bytes() == original_bytes
    output.upsert('synthetic-new', importance_raw='B', source_values={ResearchSource.IC_NET: '0'})
    upgraded = load_workbook(path)
    try:
        sheet = upgraded.active
        columns = {cell.value: cell.column for cell in sheet[1]}
        assert tuple(columns) == CANONICAL_HEADERS
        for header, value in values.items():
            assert sheet.cell(2, columns[header]).value == value
        assert sheet.cell(2, columns['INSO']).font.bold
        assert sheet.cell(2, columns['INSO']).comment.text == 'synthetic-note'
        assert sheet.cell(2, columns['IC.net']).value is None
        assert sheet.cell(3, columns['IC.net']).value == '0'
        for header in ('_inquiry_id', '_research_status'):
            assert sheet.column_dimensions[sheet.cell(1, columns[header]).column_letter].hidden
        assert upgraded['unrelated']['A1'].value == 'preserve-me'
        assert len(output.read_history()) == 2
    finally:
        upgraded.close()


def test_age_months_and_currency_conversion_for_six_month_quotes():
    result = history('findchips', [('2', calendar_month_cutoff(NOW, 6))])
    assert result.price_candidate.normalized_rmb_price == Decimal(14)
    for age in (0, 7):
        with pytest.raises(ValueError):
            replace(result.price_candidate, age_months=age)


def test_zero_reference_does_not_win_or_block_other_prices():
    icnet = IcNetAdapter(FakeClient(IC_HTML)).search('ABC-123', None, 1).source_result
    normal = _result(ResearchSource.LCSC, SourceOutcome.SUCCESS, '9')
    result = aggregate_price_results(_complete({ResearchSource.IC_NET: icnet, ResearchSource.LCSC: normal}), 3)
    assert result.status is ResearchStatus.SUCCESS
    assert result.lowest.source is ResearchSource.LCSC
    assert result.estimated_total == Decimal(27)


def test_unpriced_recent_findchips_row_does_not_block_older_valid_quote():
    html = _page(_row('ABC', '1', '[]', NOW.isoformat()),
                 _row('ABC', '1', '[[1,"USD","2"]]', calendar_month_cutoff(NOW, 4).isoformat()))
    result = FindchipsAdapter(SimpleNamespace(fetch_first_page=lambda _: FindchipsPage(html, 'https://www.findchips.com/search/ABC', NOW)), Fx(), clock=lambda: NOW).search('ABC', 10)
    assert result.price_candidate.age_months == 4
    assert result.price_candidate.normalized_rmb_price == Decimal(14)


@pytest.mark.parametrize('source', ['bom.ai', 'inso'])
def test_six_month_usd_quotes_still_use_ecb_conversion(source):
    when = calendar_month_cutoff(NOW, 6)
    if source == 'bom.ai':
        capture = BomAiCapture((BomAiPriceRecord('ABC', Decimal(2), when, 'USD'),), 'https://www.bom.ai/search/ABC', NOW)
        result = BomAiAdapter(SimpleNamespace(fetch_price_records=lambda _: capture), fx_provider=Fx(), clock=lambda: NOW).search('ABC', 3)
    else:
        capture = InsoHistoryCapture((InsoHistoryRecord(Decimal(2), 'USD', when, 'ABC'),), 'https://synthetic.invalid/history', NOW)
        result = InsoHistoryAdapter(SimpleNamespace(fetch_history=lambda _: capture), Fx(), clock=lambda: NOW).search('ABC', 3)
    assert result.price_candidate.age_months == 6
    assert result.price_candidate.normalized_rmb_price == Decimal(14)


def test_history_windows_are_independent_between_sources():
    results = [history(source, [(str(price), calendar_month_cutoff(NOW, months))])
               for source, months, price in [('findchips', 1, 10), ('hqew', 2, 2),
                                             ('bom.ai', 3, 3), ('inso', 6, 1)]]
    assert [result.price_candidate.age_months for result in results] == [1, 2, 3, 6]
    aggregation = aggregate_price_results(_complete({result.source: result for result in results}), 11)
    assert aggregation.lowest.source is ResearchSource.INSO
    assert aggregation.estimated_total == Decimal(11)


def test_selected_quote_date_is_from_the_chosen_record_not_same_price_older_row():
    older = calendar_month_cutoff(NOW, 4)
    result = history('bom.ai', [('2', older), ('2', NOW)])
    fields = {field.key: field.value for field in result.evidence.fields}
    assert result.price_candidate.age_months == 1
    assert fields['selected_observed_at'] == NOW
