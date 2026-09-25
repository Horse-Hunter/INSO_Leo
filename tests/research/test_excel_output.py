from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from src.research.excel_output import (
    CANONICAL_HEADERS,
    ESTIMATED_TOTAL_HEADER,
    IMPORTANCE_HEADER,
    INQUIRY_ID_HEADER,
    PROCESSED_AT_HEADER,
    ExcelConsistencyError,
    ExcelWriteError,
    ResearchExcelOutput,
)
from src.research.source_contracts import ResearchSource


def _columns(path: Path) -> dict[str, int]:
    worksheet = load_workbook(path).active
    return {
        worksheet.cell(1, column).value: column
        for column in range(1, worksheet.max_column + 1)
    }


def test_upsert_is_idempotent_hides_identity_and_preserves_raw_importance(
    tmp_path: Path,
) -> None:
    path = tmp_path / "调研价格.xlsx"
    output = ResearchExcelOutput(path)

    output.upsert("inq_1", importance_raw="A", remarks="first")
    output.upsert("inq_1", importance_raw="D", remarks="updated")

    worksheet = load_workbook(path).active
    columns = _columns(path)
    assert worksheet.max_row == 2
    assert worksheet.cell(2, columns[IMPORTANCE_HEADER]).value == "D"
    assert worksheet.cell(2, columns["备注"]).value == "updated"
    identity_letter = worksheet.cell(1, columns[INQUIRY_ID_HEADER]).column_letter
    assert worksheet.column_dimensions[identity_letter].hidden is True


@pytest.mark.parametrize("importance_raw", ["A", "B", "C", "D", "", None])
def test_importance_is_written_verbatim(
    tmp_path: Path, importance_raw: str | None
) -> None:
    path = tmp_path / "调研价格.xlsx"
    ResearchExcelOutput(path).upsert("inq_1", importance_raw=importance_raw)

    worksheet = load_workbook(path).active
    assert worksheet.cell(2, _columns(path)[IMPORTANCE_HEADER]).value == (
        importance_raw or None
    )


def test_previous_eight_visible_column_schema_migrates_without_fabricating_sources(
    tmp_path: Path,
) -> None:
    path = tmp_path / "调研价格.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(
        [
            "型号",
            "品牌",
            "数量",
            "重要等级",
            "货量标识",
            "预计订单总价",
            "市场最低参考价",
            "备注",
            INQUIRY_ID_HEADER,
        ]
    )
    worksheet.append(
        ["ABC", "Acme", 10, "普通", "货多", "80", "8", "old", "inq_old"]
    )
    workbook.save(path)
    workbook.close()

    ResearchExcelOutput(path).upsert(
        "inq_old",
        importance_raw="C",
        remarks=None,
        source_values={
            source: "无结果"
            for source in (
                ResearchSource.INSO,
                ResearchSource.FINDCHIPS,
                ResearchSource.HQEW,
                ResearchSource.LCSC,
                ResearchSource.BOM_AI,
            )
        },
    )

    worksheet = load_workbook(path).active
    headers = [
        worksheet.cell(1, column).value
        for column in range(1, worksheet.max_column + 1)
    ]
    row = {header: worksheet.cell(2, index + 1).value for index, header in enumerate(headers)}
    assert headers == list(CANONICAL_HEADERS)
    assert row[INQUIRY_ID_HEADER] == "inq_old"
    assert row["型号"] == "ABC"
    assert row[ESTIMATED_TOTAL_HEADER] == "80"
    assert row["INSO"] == "无结果"
    assert row["Findchips"] == "无结果"
    assert row["备注"] is None


def test_minimal_known_schema_migrates_and_retains_row_identity(tmp_path: Path) -> None:
    path = tmp_path / "调研价格.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append([INQUIRY_ID_HEADER, IMPORTANCE_HEADER, "备注"])
    worksheet.append(["inq_legacy", "普通", "legacy"])
    workbook.save(path)
    workbook.close()

    ResearchExcelOutput(path).upsert(
        "inq_legacy", importance_raw="B", mpn="ABC-1"
    )

    worksheet = load_workbook(path).active
    headers = [worksheet.cell(1, col).value for col in range(1, worksheet.max_column + 1)]
    assert headers == list(CANONICAL_HEADERS)
    assert worksheet.max_row == 2
    assert worksheet.cell(2, headers.index(INQUIRY_ID_HEADER) + 1).value == "inq_legacy"
    assert worksheet.cell(2, headers.index("型号") + 1).value == "ABC-1"


def test_duplicate_identity_anywhere_fails_closed(tmp_path: Path) -> None:
    path = tmp_path / "调研价格.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(CANONICAL_HEADERS)
    worksheet.append([None] * (len(CANONICAL_HEADERS) - 1) + ["duplicate"])
    worksheet.append([None] * (len(CANONICAL_HEADERS) - 1) + ["duplicate"])
    workbook.save(path)
    workbook.close()

    with pytest.raises(ExcelConsistencyError, match="duplicate"):
        ResearchExcelOutput(path).upsert("new", importance_raw="A")
    with pytest.raises(ExcelConsistencyError, match="duplicate"):
        ResearchExcelOutput(path).read_history()


def test_unknown_or_ambiguous_schema_fails_closed(tmp_path: Path) -> None:
    path = tmp_path / "调研价格.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append([INQUIRY_ID_HEADER, "未知列"])
    worksheet.append(["inq_1", "do not guess"])
    workbook.save(path)
    workbook.close()

    with pytest.raises(ExcelConsistencyError, match="unrecognized"):
        ResearchExcelOutput(path).upsert("inq_1", importance_raw="A")
    assert load_workbook(path).active.cell(2, 2).value == "do not guess"


def test_full_snapshot_explicit_none_clears_stale_cells(tmp_path: Path) -> None:
    path = tmp_path / "调研价格.xlsx"
    output = ResearchExcelOutput(path)
    sources = {
        source: "8"
        for source in (
            ResearchSource.INSO,
            ResearchSource.FINDCHIPS,
            ResearchSource.HQEW,
            ResearchSource.LCSC,
            ResearchSource.BOM_AI,
        )
    }
    output.upsert(
        "inq_1",
        importance_raw="A",
        mpn="ABC",
        brand="Acme",
        quantity=10,
        stock_label="货少",
        estimated_total=Decimal(80),
        market_reference="8",
        source_values=sources,
        remarks="old",
    )
    output.upsert(
        "inq_1",
        importance_raw="D",
        mpn=None,
        brand=None,
        quantity=None,
        stock_label=None,
        estimated_total=None,
        market_reference=None,
        source_values={source: "无结果" for source in sources},
        remarks=None,
    )

    worksheet = load_workbook(path).active
    columns = _columns(path)
    assert worksheet.max_row == 2
    for header in (
        "型号",
        "品牌",
        "数量",
        "货量标识",
        ESTIMATED_TOTAL_HEADER,
        "市场最低参考价",
        "备注",
    ):
        assert worksheet.cell(2, columns[header]).value is None
    assert worksheet.cell(2, columns["Findchips"]).value == "无结果"


def test_save_failure_is_wrapped_and_does_not_create_target(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "调研价格.xlsx"

    def fail_save(self, filename) -> None:
        raise RuntimeError("synthetic save failure")

    monkeypatch.setattr(Workbook, "save", fail_save)
    with pytest.raises(ExcelWriteError, match="unable to save"):
        ResearchExcelOutput(path).upsert("inq_1", importance_raw="A")
    assert not path.exists()


def test_processed_at_is_written_as_utc_iso8601_and_history_is_sorted(tmp_path: Path):
    path = tmp_path / "research.xlsx"
    output = ResearchExcelOutput(path)
    output.upsert(
        "older", importance_raw="A",
        processed_at=datetime(2026, 1, 1, 8, tzinfo=timezone.utc),
        mpn="OLD",
    )
    output.upsert(
        "newer", importance_raw="D",
        processed_at=datetime(2026, 1, 2, 8, tzinfo=timezone.utc),
        mpn="NEW",
    )
    sheet = load_workbook(path, data_only=True).active
    cols = _columns(path)
    assert sheet.cell(2, cols[PROCESSED_AT_HEADER]).value == "2026-01-01T08:00:00+00:00"
    history = output.read_history()
    assert [row.inquiry_id for row in history] == ["newer", "older"]
    assert history[0].processed_at == datetime(2026, 1, 2, 8, tzinfo=timezone.utc)


def test_legacy_history_has_no_inferred_time_and_migrates_without_losing_values(tmp_path: Path):
    path = tmp_path / "legacy.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([INQUIRY_ID_HEADER, IMPORTANCE_HEADER, "备注"])
    sheet.append(["inq_legacy", "C", "keep this"])
    workbook.save(path)
    workbook.close()

    output = ResearchExcelOutput(path)
    legacy = output.read_history()
    assert len(legacy) == 1 and legacy[0].processed_at is None
    output.upsert("inq_new", importance_raw="A", processed_at=datetime(2026, 3, 1, tzinfo=timezone.utc))
    migrated = output.read_history()
    old = next(row for row in migrated if row.inquiry_id == "inq_legacy")
    assert old.processed_at is None
    assert old.importance_raw == "C"
    assert old.remarks == "keep this"


def test_history_keeps_legacy_rows_stable_after_timestamped_rows(tmp_path: Path):
    path = tmp_path / "legacy-order.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([INQUIRY_ID_HEADER, IMPORTANCE_HEADER, "备注"])
    sheet.append(["legacy-first", "A", "first"])
    sheet.append(["legacy-second", "B", "second"])
    sheet.append(["dated", "D", "dated"])
    workbook.save(path)
    workbook.close()

    output = ResearchExcelOutput(path)
    output.upsert(
        "dated", importance_raw="D",
        processed_at=datetime(2026, 7, 1, tzinfo=timezone.utc),
    )
    history = output.read_history()
    assert [row.inquiry_id for row in history] == [
        "dated", "legacy-first", "legacy-second"
    ]
    assert [row.processed_at for row in history[1:]] == [None, None]
