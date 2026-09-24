import json
import socket
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.core import CredentialNotConfiguredError, Login
from src.research.bom_ai import BomAiBrowserConfig, BomAiClientError
from src.research.contracts import ResearchInput, ResearchStatus
from src.research.credentials import CoreResearchCredentials
from src.research.excel_output import INQUIRY_ID_HEADER
from src.research.findchips import FindchipsPageUnavailable
from src.research.fx import UsdRmbQuote
from src.research.hqew import HqewPageUnavailable
from src.research.icnet import (
    CdpIcNetClient,
    IcNetPageUnavailable,
    PlaywrightIcNetClient,
)
from src.research.inso_history import INSO_SITE_ID, InsoBrowserConfig, InsoReadError
from src.research.lcsc import LcscPageUnavailable
from src.research.runtime import (
    BrowserRuntimeConfig,
    CdpRuntimeConfig,
    IcNetRuntimeConfig,
    ResearchRuntimeConfig,
    ResearchRuntimeConfigError,
    ResearchRuntimePrerequisiteError,
    assess_readiness,
    build_icnet_client,
    build_production_research_service,
    build_research_service,
    format_readiness,
    load_runtime_config,
    probe_loopback_endpoint,
)
from src.research.service import ResearchService

NOW = datetime(2026, 9, 24, 10, tzinfo=UTC)


def _mapping(**overrides: object) -> dict:
    values: dict = {
        "excel_output_path": "data/调研价格.xlsx",
        "browser": {
            "channel": "chrome",
            "headless": False,
            "timeout_ms": 45_000,
            "settle_ms": 3_000,
        },
        "cdp": {"cdp_url": "http://127.0.0.1:9222"},
        "icnet": {"mode": "credentials"},
        "bom_ai": {
            "login_url": "https://www.bom.ai/login",
            "result_url_template": "https://www.bom.ai/search/{mpn}",
            "username_selector": "#user",
            "password_selector": "#pass",
            "login_button_selector": "#signin",
        },
        "inso": {
            "login_url": "https://inso.example/login",
            "username_selector": "#u",
            "password_selector": "#p",
            "login_button_selector": "#l",
            "mpn_selector": "#m",
            "query_button_selector": "#q",
            "date_headers": ["报价时间"],
        },
    }
    values.update(overrides)
    return values


def _bom_ai_config() -> BomAiBrowserConfig:
    return BomAiBrowserConfig(**_mapping()["bom_ai"])


def _inso_config() -> InsoBrowserConfig:
    section = dict(_mapping()["inso"])
    section["date_headers"] = tuple(section["date_headers"])
    return InsoBrowserConfig(**section)


def _write_config(tmp_path: Path, mapping: dict | None = None) -> Path:
    path = tmp_path / "research.json"
    path.write_text(
        json.dumps(mapping if mapping is not None else _mapping(), ensure_ascii=False),
        encoding="utf-8",
    )
    return path


def _tmp_config(tmp_path: Path) -> ResearchRuntimeConfig:
    return load_runtime_config(
        _write_config(
            tmp_path,
            _mapping(excel_output_path=str(tmp_path / "调研价格.xlsx")),
        )
    )


class AllSitesProvider:
    """Core CredentialProvider double that authorizes every Research site."""

    def __init__(self) -> None:
        self.requested: list[str] = []

    def get_login(self, site_id: str) -> Login:
        self.requested.append(site_id)
        return Login(site_id, "", "synthetic-user", "synthetic-password", None)


class InsoMissingProvider(AllSitesProvider):
    def get_login(self, site_id: str) -> Login:
        if site_id == INSO_SITE_ID:
            raise CredentialNotConfiguredError(site_id, "password=must-not-leak")
        return super().get_login(site_id)


class FakeFx:
    def get_quote(self) -> UsdRmbQuote:
        return UsdRmbQuote(Decimal(7), NOW, "synthetic-test-only")


class RecordingIcNetClient:
    def __init__(self) -> None:
        self.queries: list[str] = []

    def fetch_first_page(self, mpn: str):
        self.queries.append(mpn)
        raise IcNetPageUnavailable("SYNTHETIC_UNAVAILABLE")


class RecordingFindchipsClient:
    def __init__(self) -> None:
        self.queries: list[str] = []

    def fetch_first_page(self, mpn: str):
        self.queries.append(mpn)
        raise FindchipsPageUnavailable("SYNTHETIC_UNAVAILABLE")


class RecordingHqewClient:
    def __init__(self) -> None:
        self.queries: list[str] = []

    def fetch_first_page(self, mpn: str):
        self.queries.append(mpn)
        raise HqewPageUnavailable("SYNTHETIC_UNAVAILABLE")


class RecordingLcscClient:
    def __init__(self) -> None:
        self.queries: list[str] = []

    def fetch_product_page(self, mpn: str):
        self.queries.append(mpn)
        raise LcscPageUnavailable("SYNTHETIC_UNAVAILABLE")


class RecordingBomAiBrowser:
    def __init__(self) -> None:
        self.queries: list[str] = []

    def fetch_price_page(self, mpn: str, login: object):
        self.queries.append(mpn)
        raise BomAiClientError("SYNTHETIC_UNAVAILABLE")


class RecordingInsoBrowser:
    def __init__(self) -> None:
        self.queries: list[str] = []

    def fetch_procurement_temporary_inquiry_history(self, mpn: str, login: object):
        self.queries.append(mpn)
        raise InsoReadError("SYNTHETIC_UNAVAILABLE")


def test_load_runtime_config_reads_non_secret_values(tmp_path: Path) -> None:
    config = load_runtime_config(_write_config(tmp_path))

    assert config.excel_output_path == Path("data/调研价格.xlsx")
    assert config.bom_ai.login_url == "https://www.bom.ai/login"
    assert config.bom_ai.result_url("ABC") == "https://www.bom.ai/search/ABC"
    assert config.inso.login_url == "https://inso.example/login"
    assert config.inso.date_headers == ("报价时间",)
    assert config.browser.timeout_ms == 45_000
    assert config.cdp.cdp_url == "http://127.0.0.1:9222"
    assert config.icnet.mode == "credentials"


def test_missing_or_invalid_configuration_fails_closed(tmp_path: Path) -> None:
    with pytest.raises(ResearchRuntimeConfigError):
        load_runtime_config(tmp_path / "absent.json")

    broken = tmp_path / "broken.json"
    broken.write_text("{not json", encoding="utf-8")
    with pytest.raises(ResearchRuntimeConfigError):
        load_runtime_config(broken)

    missing_section = _mapping()
    del missing_section["bom_ai"]
    with pytest.raises(ResearchRuntimeConfigError):
        load_runtime_config(_write_config(tmp_path, missing_section))

    missing_key = _mapping()
    del missing_key["inso"]["query_button_selector"]
    with pytest.raises(ResearchRuntimeConfigError):
        load_runtime_config(_write_config(tmp_path, missing_key))

    wrong_host = _mapping()
    wrong_host["bom_ai"]["login_url"] = "https://evil.example/login"
    with pytest.raises(ResearchRuntimeConfigError):
        load_runtime_config(_write_config(tmp_path, wrong_host))

    wrong_type = _mapping()
    wrong_type["browser"]["timeout_ms"] = "soon"
    with pytest.raises(ResearchRuntimeConfigError):
        load_runtime_config(_write_config(tmp_path, wrong_type))


def test_browser_cdp_and_icnet_runtime_validation() -> None:
    with pytest.raises(ValueError):
        BrowserRuntimeConfig(timeout_ms=0)
    with pytest.raises(ValueError):
        BrowserRuntimeConfig(settle_ms=-1)
    with pytest.raises(TypeError):
        BrowserRuntimeConfig(headless="no")  # type: ignore[arg-type]
    with pytest.raises(TypeError):
        BrowserRuntimeConfig(channel=7)  # type: ignore[arg-type]
    with pytest.raises(ValueError):
        CdpRuntimeConfig(cdp_url="http://127.0.0.1")
    with pytest.raises(ValueError):
        CdpRuntimeConfig(cdp_url="http://example.com:9222")
    with pytest.raises(ValueError):
        IcNetRuntimeConfig(mode="headless-stealth")
    assert CdpRuntimeConfig(cdp_url="http://localhost:9222").cdp_url.endswith("9222")
    assert IcNetRuntimeConfig(mode="cdp").mode == "cdp"


def test_icnet_acquisition_mode_is_selected_by_runtime_config() -> None:
    credentials = CoreResearchCredentials(AllSitesProvider())

    default_config = ResearchRuntimeConfig(
        excel_output_path=Path("out.xlsx"),
        bom_ai=_bom_ai_config(),
        inso=_inso_config(),
    )
    cdp_config = ResearchRuntimeConfig(
        excel_output_path=Path("out.xlsx"),
        bom_ai=_bom_ai_config(),
        inso=_inso_config(),
        icnet=IcNetRuntimeConfig(mode="cdp"),
        cdp=CdpRuntimeConfig(cdp_url="http://127.0.0.1:9333"),
    )

    assert isinstance(
        build_icnet_client(default_config, credentials), PlaywrightIcNetClient
    )
    assert isinstance(build_icnet_client(cdp_config, credentials), CdpIcNetClient)


def test_invalid_icnet_mode_in_config_fails_closed(tmp_path: Path) -> None:
    mapping = _mapping()
    mapping["icnet"] = {"mode": "stealth"}

    with pytest.raises(ResearchRuntimeConfigError):
        load_runtime_config(_write_config(tmp_path, mapping))


def test_composition_calls_every_canonical_source_and_keeps_excel_idempotent(
    tmp_path: Path,
) -> None:
    config = _tmp_config(tmp_path)
    icnet = RecordingIcNetClient()
    findchips = RecordingFindchipsClient()
    hqew = RecordingHqewClient()
    lcsc = RecordingLcscClient()
    bom_ai = RecordingBomAiBrowser()
    inso = RecordingInsoBrowser()

    service = build_research_service(
        config,
        provider=AllSitesProvider(),
        fx_provider=FakeFx(),
        clock=lambda: NOW,
        icnet_client=icnet,
        findchips_client=findchips,
        hqew_client=hqew,
        lcsc_client=lcsc,
        bom_ai_browser=bom_ai,
        inso_browser=inso,
    )
    research_input = ResearchInput("inq_runtime", "ABC-1", None, 10, "A")

    assert service.execute(research_input).status is ResearchStatus.RETRYABLE_FAILURE
    assert icnet.queries == ["ABC-1"]
    for source in (findchips, hqew, lcsc, bom_ai, inso):
        assert source.queries == ["ABC-1"]

    service.execute(research_input)
    worksheet = load_workbook(config.excel_output_path).active
    headers = [
        worksheet.cell(1, column).value
        for column in range(1, worksheet.max_column + 1)
    ]
    assert INQUIRY_ID_HEADER in headers
    assert worksheet.max_row == 2
    assert worksheet.cell(2, headers.index(INQUIRY_ID_HEADER) + 1).value == (
        "inq_runtime"
    )


def test_readiness_reports_missing_site_and_manual_action(tmp_path: Path) -> None:
    config = _tmp_config(tmp_path)

    readiness = assess_readiness(
        config.cdp.cdp_url,
        provider=InsoMissingProvider(),
        cdp_probe=lambda url: False,
    )

    assert readiness.missing_site_ids == (INSO_SITE_ID,)
    assert readiness.ready is False
    issues = readiness.issues()
    assert any(INSO_SITE_ID in issue for issue in issues)
    assert any("CDP" in issue for issue in issues)
    action = readiness.manual_action()
    assert action is not None and INSO_SITE_ID in action and "Chrome" in action

    report = format_readiness(readiness)
    assert "runtime readiness: BLOCKED" in report
    assert "must-not-leak" not in report
    assert "synthetic-password" not in report


def test_readiness_does_not_require_dom_configuration() -> None:
    readiness = assess_readiness(
        "http://127.0.0.1:9222",
        provider=AllSitesProvider(),
        cdp_probe=lambda url: url == "http://127.0.0.1:9222",
    )

    assert readiness.ready is True
    assert readiness.missing_site_ids == ()


def test_ready_configuration_composes_a_production_service(tmp_path: Path) -> None:
    config = _tmp_config(tmp_path)

    readiness = assess_readiness(
        config.cdp.cdp_url,
        provider=AllSitesProvider(),
        cdp_probe=lambda url: True,
    )
    assert readiness.ready is True
    assert readiness.manual_action() is None
    assert "runtime readiness: READY" in format_readiness(readiness)

    service = build_production_research_service(
        config, provider=AllSitesProvider(), cdp_probe=lambda url: True
    )
    assert isinstance(service, ResearchService)


def test_not_ready_configuration_fails_closed(tmp_path: Path) -> None:
    config = _tmp_config(tmp_path)

    with pytest.raises(ResearchRuntimePrerequisiteError) as error:
        build_production_research_service(
            config,
            provider=InsoMissingProvider(),
            cdp_probe=lambda url: True,
        )

    assert INSO_SITE_ID in str(error.value)


def test_probe_loopback_endpoint_detects_a_local_listener() -> None:
    listener = socket.socket()
    listener.bind(("127.0.0.1", 0))
    listener.listen(1)
    port = listener.getsockname()[1]
    try:
        assert probe_loopback_endpoint(f"http://127.0.0.1:{port}") is True
    finally:
        listener.close()

    assert probe_loopback_endpoint(f"http://127.0.0.1:{port}") is False
    with pytest.raises(ValueError):
        probe_loopback_endpoint("http://example.com:9222")
