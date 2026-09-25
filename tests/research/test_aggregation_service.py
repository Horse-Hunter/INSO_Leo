from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from src.research.aggregation import aggregate_price_results
from src.research.contracts import ResearchInput, ResearchReasonCode, ResearchStatus
from src.research.excel_output import (
    CANONICAL_HEADERS,
    PROCESSED_AT_HEADER,
    RESEARCH_STATUS_HEADER,
    ExcelWriteError,
    ResearchExcelOutput,
)
from src.research.icnet import IcNetResult
from src.research.service import ResearchService
from src.research.source_contracts import (
    PRICE_SOURCES,
    EvidenceField,
    PriceCandidate,
    ResearchSource,
    SourceEvidence,
    SourceOutcome,
    SourceResult,
)

NOW = datetime(2026, 9, 22, tzinfo=UTC)


def _result(
    source: ResearchSource,
    outcome: SourceOutcome,
    price: str | None = None,
    *,
    out_of_stock: str | None = None,
    failure_code: str = "HTTP_REQUEST_FAILED",
) -> SourceResult:
    normal = (
        PriceCandidate(source, "ABC", Decimal(price), "RMB", Decimal(price), NOW)
        if price is not None
        else None
    )
    fallback = (
        PriceCandidate(
            source,
            "ABC",
            Decimal(out_of_stock),
            "RMB",
            Decimal(out_of_stock),
            NOW,
        )
        if out_of_stock is not None
        else None
    )
    evidence = SourceEvidence(
        source,
        "ABC",
        "ABC" if normal or fallback else None,
        outcome,
        NOW,
        fields=(
            (EvidenceField("failure_code", failure_code),)
            if outcome is SourceOutcome.SOURCE_UNAVAILABLE
            else ()
        ),
    )
    return SourceResult(source, outcome, evidence, normal, fallback)


def _complete(
    overrides: dict[ResearchSource, SourceResult] | None = None,
) -> tuple[SourceResult, ...]:
    values = {
        source: _result(source, SourceOutcome.NO_VALID_PRICE)
        for source in PRICE_SOURCES
    }
    values.update(overrides or {})
    return tuple(values[source] for source in PRICE_SOURCES)


def test_normal_pool_20_percent_and_partial_failure() -> None:
    results = _complete(
        {
            ResearchSource.FINDCHIPS: _result(
                ResearchSource.FINDCHIPS,
                SourceOutcome.SUCCESS,
                "8",
                out_of_stock="1",
            ),
            ResearchSource.HQEW: _result(
                ResearchSource.HQEW, SourceOutcome.SUCCESS, "11"
            ),
            ResearchSource.BOM_AI: _result(
                ResearchSource.BOM_AI,
                SourceOutcome.SOURCE_UNAVAILABLE,
                failure_code="RESULT_ROWS_MISSING",
            ),
        }
    )

    aggregation = aggregate_price_results(results, 10)

    assert aggregation.status is ResearchStatus.PARTIAL_SUCCESS
    assert aggregation.market_reference == "8\n11-华强"
    assert aggregation.estimated_total == Decimal(80)
    assert aggregation.used_out_of_stock_fallback is False
    assert aggregation.remarks == "正能量：结果解析失败"
    assert [candidate.normalized_rmb_price for candidate in aggregation.candidates] == [
        Decimal(8),
        Decimal(11),
    ]


def test_http_forbidden_is_not_reported_as_a_parse_failure() -> None:
    results = _complete(
        {
            ResearchSource.HQEW: _result(
                ResearchSource.HQEW,
                SourceOutcome.SOURCE_UNAVAILABLE,
                failure_code="HTTP_STATUS_403",
            ),
        }
    )

    aggregation = aggregate_price_results(results, 10)

    assert aggregation.remarks == "华强：访问被站点拒绝"


def test_no_normal_price_with_technical_failure_is_retryable_not_fallback() -> None:
    results = _complete(
        {
            ResearchSource.FINDCHIPS: _result(
                ResearchSource.FINDCHIPS,
                SourceOutcome.SUCCESS,
                out_of_stock="2",
            ),
            ResearchSource.HQEW: _result(
                ResearchSource.HQEW,
                SourceOutcome.SOURCE_UNAVAILABLE,
                failure_code="INTERACTIVE_CHALLENGE_REQUIRED",
            ),
        }
    )

    aggregation = aggregate_price_results(results, 10)

    assert aggregation.status is ResearchStatus.RETRYABLE_FAILURE
    assert aggregation.market_reference is None
    assert aggregation.estimated_total is None
    assert aggregation.remarks == "华强：需要人工验证"


def test_multiple_failure_remarks_follow_canonical_source_order() -> None:
    results = list(
        _complete(
            {
                ResearchSource.FINDCHIPS: _result(
                    ResearchSource.FINDCHIPS,
                    SourceOutcome.SOURCE_UNAVAILABLE,
                    failure_code="HTTP_REQUEST_FAILED",
                ),
                ResearchSource.INSO: _result(
                    ResearchSource.INSO,
                    SourceOutcome.SOURCE_UNAVAILABLE,
                    failure_code="CREDENTIALS_UNAVAILABLE",
                ),
            }
        )
    )
    results.reverse()
    aggregation = aggregate_price_results(tuple(results), 1)
    assert aggregation.remarks == "Findchips：暂时不可用；INSO：登录不可用"


def test_authenticated_session_failure_is_reported_as_login_unavailable() -> None:
    results = _complete(
        {
            ResearchSource.LCSC: _result(
                ResearchSource.LCSC,
                SourceOutcome.SOURCE_UNAVAILABLE,
                failure_code="AUTHENTICATED_SESSION_REQUIRED",
            ),
        }
    )
    aggregation = aggregate_price_results(tuple(results), 1)
    assert aggregation.remarks == "立创：登录不可用"


def test_out_of_stock_fallback_is_partial_and_uses_20_percent_rule() -> None:
    results = _complete(
        {
            ResearchSource.FINDCHIPS: _result(
                ResearchSource.FINDCHIPS,
                SourceOutcome.SUCCESS,
                out_of_stock="8",
            ),
            ResearchSource.LCSC: _result(
                ResearchSource.LCSC,
                SourceOutcome.SUCCESS,
                out_of_stock="11",
            ),
        }
    )

    aggregation = aggregate_price_results(results, 3)

    assert aggregation.status is ResearchStatus.PARTIAL_SUCCESS
    assert aggregation.used_out_of_stock_fallback is True
    assert aggregation.market_reference == "8\n11-立创"
    assert aggregation.estimated_total == Decimal(24)
    assert aggregation.remarks == "仅找到无库存报价"


def test_no_price_and_no_technical_failure_is_terminal_exception() -> None:
    aggregation = aggregate_price_results(_complete(), 1)

    assert aggregation.status is ResearchStatus.EXCEPTION
    assert aggregation.reason_code is ResearchReasonCode.NO_MATCHING_PRODUCT
    assert aggregation.market_reference is None
    assert aggregation.remarks == "五个价格来源均无报价，可能是客户填写的型号有误"


def test_no_quote_exception_is_persisted_in_research_history(tmp_path: Path) -> None:
    path = tmp_path / "调研价格.xlsx"
    result = _service(path, _complete()).execute(
        ResearchInput("inq_no_quote", "UNKNOWN-MPN", None, 4, "B")
    )

    worksheet = load_workbook(path).active
    headers = {
        worksheet.cell(1, column).value: column
        for column in range(1, worksheet.max_column + 1)
    }
    assert result.status is ResearchStatus.EXCEPTION
    assert worksheet.cell(2, headers[RESEARCH_STATUS_HEADER]).value == "EXCEPTION"
    assert "可能是客户填写的型号有误" in worksheet.cell(2, headers["备注"]).value


class FakeIcNet:
    def search(
        self, target_mpn: str, input_brand: str | None, customer_quantity: int
    ) -> IcNetResult:
        evidence = SourceEvidence(
            ResearchSource.IC_NET,
            target_mpn,
            target_mpn,
            SourceOutcome.SUCCESS,
            NOW,
        )
        return IcNetResult(
            SourceResult(ResearchSource.IC_NET, SourceOutcome.SUCCESS, evidence),
            input_brand or "Acme",
            "货多",
        )


class FakePrice:
    def __init__(self, result: SourceResult) -> None:
        self.result = result

    def search(self, target_mpn: str, customer_quantity: int) -> SourceResult:
        return self.result


def _service(path: Path, values: tuple[SourceResult, ...]) -> ResearchService:
    by_source = {value.source: value for value in values}
    return ResearchService(
        icnet=FakeIcNet(),
        findchips=FakePrice(by_source[ResearchSource.FINDCHIPS]),
        hqew=FakePrice(by_source[ResearchSource.HQEW]),
        lcsc=FakePrice(by_source[ResearchSource.LCSC]),
        bom_ai=FakePrice(by_source[ResearchSource.BOM_AI]),
        inso=FakePrice(by_source[ResearchSource.INSO]),
        output=ResearchExcelOutput(path),
    )


def test_icnet_challenge_is_visible_in_excel_remarks(tmp_path: Path) -> None:
    class ChallengedIcNet:
        def search(
            self, target_mpn: str, input_brand: str | None, customer_quantity: int
        ) -> IcNetResult:
            source = _result(
                ResearchSource.IC_NET,
                SourceOutcome.SOURCE_UNAVAILABLE,
                failure_code="INTERACTIVE_CHALLENGE_REQUIRED",
            )
            return IcNetResult(source)

    path = tmp_path / "price.xlsx"
    service = _service(path, _complete())
    service._icnet = ChallengedIcNet()
    result = service.execute(ResearchInput("inq_challenge", "ABC", None, 10, "A"))
    assert result.status is ResearchStatus.EXCEPTION
    assert result.remarks == (
        "五个价格来源均无报价，可能是客户填写的型号有误；IC.net：需要人工验证"
    )
    assert load_workbook(path).active["M2"].value == result.remarks


def test_service_writes_canonical_idempotent_full_snapshot(tmp_path: Path) -> None:
    path = tmp_path / "调研价格.xlsx"
    sources = _complete(
        {
            ResearchSource.INSO: _result(
                ResearchSource.INSO, SourceOutcome.SUCCESS, "9"
            ),
            ResearchSource.FINDCHIPS: _result(
                ResearchSource.FINDCHIPS,
                SourceOutcome.SUCCESS,
                "8",
                out_of_stock="7",
            ),
        }
    )
    value = ResearchInput("inq_1", "ABC", None, 10, "A")

    assert _service(path, sources).execute(value).status is ResearchStatus.SUCCESS
    assert _service(path, sources).execute(value).status is ResearchStatus.SUCCESS

    worksheet = load_workbook(path).active
    headers = [worksheet.cell(1, col).value for col in range(1, worksheet.max_column + 1)]
    row = {
        headers[index - 1]: worksheet.cell(2, index).value
        for index in range(1, worksheet.max_column + 1)
    }
    assert headers == [*CANONICAL_HEADERS]
    assert worksheet.max_row == 2
    assert row["重要等级"] == "A"
    assert row["预估订单总价"] == "80"
    assert row["INSO"] == "9"
    assert row["Findchips"] == "8\n7（无库存）"
    assert row["华强"] == "无结果"
    assert row["备注"] is None
    assert row[RESEARCH_STATUS_HEADER] == ResearchStatus.SUCCESS.value
    assert row[PROCESSED_AT_HEADER].endswith("+00:00")
    inquiry_letter = worksheet.cell(1, worksheet.max_column).column_letter
    assert worksheet.column_dimensions[inquiry_letter].hidden is True


def test_retry_snapshot_clears_stale_prices_sources_and_remarks(tmp_path: Path) -> None:
    path = tmp_path / "调研价格.xlsx"
    value = ResearchInput("inq_refresh", "ABC", None, 10, "C")
    success = _complete(
        {
            ResearchSource.FINDCHIPS: _result(
                ResearchSource.FINDCHIPS, SourceOutcome.SUCCESS, "8"
            )
        }
    )
    retry = _complete(
        {
            ResearchSource.FINDCHIPS: _result(
                ResearchSource.FINDCHIPS, SourceOutcome.NO_VALID_PRICE
            ),
            ResearchSource.INSO: _result(
                ResearchSource.INSO,
                SourceOutcome.SOURCE_UNAVAILABLE,
                failure_code="BROWSER_TIMEOUT",
            ),
        }
    )

    _service(path, success).execute(value)
    result = _service(path, retry).execute(value)

    worksheet = load_workbook(path).active
    columns = {
        worksheet.cell(1, col).value: col
        for col in range(1, worksheet.max_column + 1)
    }
    assert result.status is ResearchStatus.RETRYABLE_FAILURE
    assert worksheet.max_row == 2
    assert worksheet.cell(2, columns["预估订单总价"]).value is None
    assert worksheet.cell(2, columns["市场最低参考价"]).value is None
    assert worksheet.cell(2, columns["Findchips"]).value == "无结果"
    assert worksheet.cell(2, columns["INSO"]).value == "无结果"
    assert worksheet.cell(2, columns["备注"]).value == "INSO：暂时不可用"


def test_excel_failure_is_retryable_and_keeps_acquisition_evidence() -> None:
    class FailingOutput:
        def upsert(self, inquiry_id: str, **values: object) -> None:
            raise ExcelWriteError("synthetic failure")

    values = _complete(
        {
            ResearchSource.FINDCHIPS: _result(
                ResearchSource.FINDCHIPS, SourceOutcome.SUCCESS, "8"
            )
        }
    )
    by_source = {value.source: value for value in values}
    service = ResearchService(
        icnet=FakeIcNet(),
        findchips=FakePrice(by_source[ResearchSource.FINDCHIPS]),
        hqew=FakePrice(by_source[ResearchSource.HQEW]),
        lcsc=FakePrice(by_source[ResearchSource.LCSC]),
        bom_ai=FakePrice(by_source[ResearchSource.BOM_AI]),
        inso=FakePrice(by_source[ResearchSource.INSO]),
        output=FailingOutput(),  # type: ignore[arg-type]
    )

    execution = service.execute_detailed(
        ResearchInput("inq_failure", "ABC", None, 10, "D")
    )

    assert execution.result.status is ResearchStatus.RETRYABLE_FAILURE
    assert execution.price_sources == values
    assert execution.aggregation.lowest is not None


def test_corrupted_workbook_is_retryable(tmp_path: Path) -> None:
    path = tmp_path / "调研价格.xlsx"
    path.write_bytes(b"not an xlsx archive")
    values = _complete(
        {
            ResearchSource.FINDCHIPS: _result(
                ResearchSource.FINDCHIPS, SourceOutcome.SUCCESS, "8"
            )
        }
    )

    result = _service(path, values).execute(
        ResearchInput("inq_corrupt", "ABC", None, 10, "A")
    )

    assert result.status is ResearchStatus.RETRYABLE_FAILURE
