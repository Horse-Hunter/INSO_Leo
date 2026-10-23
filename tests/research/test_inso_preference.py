"""Monthly INSO selection and precision, using synthetic source/Excel data only."""

from dataclasses import replace
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from src.research.aggregation import aggregate_price_results
from src.research.contracts import ResearchInput, ResearchStatus
from src.research.excel_output import CANONICAL_HEADERS, ResearchExcelOutput
from src.research.icnet import IcNetAdapter
from src.research.service import ResearchService
from src.research.source_contracts import (
    PRICE_SOURCES,
    ResearchSource,
    SourceOutcome,
    format_source_result,
    money_text,
)
from tests.research.test_aggregation_service import FakePrice, _complete, _result
from tests.research.test_icnet import FakeClient
from tests.research.test_six_price_sources import IC_HTML, reference


def inso_price(value, months):
    result = _result(ResearchSource.INSO, SourceOutcome.SUCCESS, value)
    return replace(result, price_candidate=replace(result.price_candidate, age_months=months))


@pytest.mark.parametrize('months,limit', [(1, '11.2'), (2, '11'), (3, '10.8'),
    (4, '10.6'), (5, '10.4'), (6, '10.2')])
@pytest.mark.parametrize('difference', ['-0.000000001', '0', '0.000000001'])
def test_inclusive_tolerance_boundaries_use_unrounded_decimal(months, limit, difference):
    minimum = _result(ResearchSource.IC_NET, SourceOutcome.SUCCESS, '10')
    price = Decimal(limit) + Decimal(difference)
    inso = inso_price(str(price), months)
    sources = _complete({ResearchSource.IC_NET: minimum, ResearchSource.INSO: inso})
    original = tuple(result.price_candidate for result in sources)
    result = aggregate_price_results(sources, 137)
    chosen = price if Decimal(difference) <= 0 else Decimal(10)
    assert result.estimated_total == chosen * Decimal(137)
    assert result.market_reference == money_text(chosen)
    assert '\n' not in result.market_reference
    assert result.lowest is minimum.price_candidate  # Still the actual six-source minimum.
    assert tuple(item.price_candidate for item in sources) == original
    assert inso.price_candidate.raw_price == price
    assert result.status is ResearchStatus.SUCCESS


@pytest.mark.parametrize('months', range(1, 7))
def test_inso_itself_lowest_or_only_valid_price(months):
    inso = inso_price('2.345', months)
    for overrides in ({ResearchSource.INSO: inso}, {ResearchSource.INSO: inso,
                      ResearchSource.HQEW: _result(ResearchSource.HQEW, SourceOutcome.SUCCESS, '9')}):
        result = aggregate_price_results(_complete(overrides), 3)
        assert result.lowest is inso.price_candidate
        assert result.estimated_total == Decimal('7.035')
        assert result.market_reference == money_text(Decimal('2.345'))


@pytest.mark.parametrize('outcome', [SourceOutcome.NO_VALID_PRICE,
    SourceOutcome.NO_STRICT_MPN_MATCH, SourceOutcome.SOURCE_UNAVAILABLE])
def test_no_valid_inso_price_uses_actual_minimum(outcome):
    sources = _complete({ResearchSource.IC_NET: _result(ResearchSource.IC_NET, SourceOutcome.SUCCESS, '0.001'),
        ResearchSource.INSO: _result(ResearchSource.INSO, outcome)})
    result = aggregate_price_results(sources, 3)
    assert result.market_reference == '0.001'
    assert result.estimated_total == Decimal('0.003')
    assert result.status is (ResearchStatus.PARTIAL_SUCCESS if outcome is SourceOutcome.SOURCE_UNAVAILABLE else ResearchStatus.SUCCESS)


@pytest.mark.parametrize('raw,expected', [('0.001', '0.001'), ('0.00001', '0.00001'),
    ('1E-30', '0.' + '0' * 29 + '1'), ('0.004999', '0.004999'), ('0.005', '0.01'),
    ('0', '0'), ('0.0000', '0'), ('12.345', '12.35')])
def test_display_never_erases_a_positive_price(raw, expected):
    value = Decimal(raw)
    assert money_text(value) == expected
    assert value == Decimal(raw)
    if value > 0:
        assert Decimal(money_text(value)) > 0


def test_preferred_micro_price_and_total_keep_calculation_precision():
    sources = _complete({ResearchSource.IC_NET: _result(ResearchSource.IC_NET, SourceOutcome.SUCCESS, '0.001'),
                         ResearchSource.INSO: inso_price('0.00112', 1)})
    result = aggregate_price_results(sources, 3)
    assert result.lowest.normalized_rmb_price == Decimal('0.001')
    assert result.market_reference == '0.00112'
    assert result.estimated_total == Decimal('0.00336')
    assert format_source_result(sources[0]) == '0.001'


@pytest.mark.parametrize('inso', [None, '0.00112', '0.0011200000001'])
def test_real_icnet_capture_to_existing_excel_fields_preserves_subcent_prices(tmp_path, inso):
    client = FakeClient(IC_HTML.replace('<body>', '<body>' + reference('0.001')))
    values = _complete({ResearchSource.INSO: inso_price(inso, 1)} if inso else {})
    by_source = {item.source: item for item in values}
    path = tmp_path / 'synthetic-prices.xlsx'
    service = ResearchService(icnet=IcNetAdapter(client),
        findchips=FakePrice(by_source[ResearchSource.FINDCHIPS]), hqew=FakePrice(by_source[ResearchSource.HQEW]),
        lcsc=FakePrice(by_source[ResearchSource.LCSC]), bom_ai=FakePrice(by_source[ResearchSource.BOM_AI]),
        inso=FakePrice(by_source[ResearchSource.INSO]), output=ResearchExcelOutput(path))
    executed = service.execute_detailed(ResearchInput('synthetic-precision', 'ABC-123', None, 3, 'A'))
    chosen = Decimal('0.00112') if inso == '0.00112' else Decimal('0.001')
    assert executed.aggregation.estimated_total == chosen * 3
    assert client.queries == ['ABC-123']
    wb = load_workbook(path)
    try:
        ws = wb.active
        columns = {cell.value: cell.column for cell in ws[1]}
        assert tuple(columns) == CANONICAL_HEADERS  # No new columns.
        assert tuple(columns) == ('型号', '品牌', '数量', '重要等级', '货量标识',
            '预估订单总价', '市场最低参考价', 'INSO', 'Findchips', '华强', '立创',
            '正能量', 'IC.net', '备注', '处理时间', '_research_status', '_inquiry_id')
        assert ws.cell(2, columns['IC.net']).value == '0.001'
        assert ws.cell(2, columns['市场最低参考价']).value == money_text(chosen)
        assert ws.cell(2, columns['预估订单总价']).value == money_text(chosen * 3)
        if inso:
            assert ws.cell(2, columns['INSO']).value == money_text(Decimal(inso))
    finally:
        wb.close()


def test_true_zero_is_still_excluded_and_cannot_trigger_inso_preference():
    client = FakeClient(IC_HTML.replace('<body>', '<body>' + reference('0')))
    icnet = IcNetAdapter(client).search('ABC-123', None, 10).source_result
    assert icnet.price_candidate is None
    assert format_source_result(icnet) == '0'
    result = aggregate_price_results(_complete({ResearchSource.IC_NET: icnet, ResearchSource.INSO: inso_price('0.001', 6)}), 10)
    assert result.market_reference == '0.001'
    assert result.estimated_total == Decimal('0.01')


def test_preference_keeps_failure_and_fallback_rules_and_raw_sources():
    inso = inso_price('11.2', 1)
    sources = _complete({ResearchSource.IC_NET: _result(ResearchSource.IC_NET, SourceOutcome.SUCCESS, '10'),
        ResearchSource.INSO: inso, ResearchSource.HQEW: _result(ResearchSource.HQEW, SourceOutcome.SOURCE_UNAVAILABLE)})
    original_cells = tuple(format_source_result(item) for item in sources)
    result = aggregate_price_results(sources, 2)
    assert result.status is ResearchStatus.PARTIAL_SUCCESS
    assert result.estimated_total == Decimal('22.4')
    assert tuple(format_source_result(item) for item in sources) == original_cells
    fallback = aggregate_price_results(_complete({ResearchSource.FINDCHIPS:
        _result(ResearchSource.FINDCHIPS, SourceOutcome.SUCCESS, out_of_stock='0.001')}), 2)
    assert fallback.used_out_of_stock_fallback
    assert fallback.market_reference == '0.001'
    assert fallback.estimated_total == Decimal('0.002')
    assert len(PRICE_SOURCES) == 6
