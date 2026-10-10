from datetime import date
from decimal import Decimal
from email.message import EmailMessage
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from src.order_mail.contracts import (
    OrderError,
    delivery_date,
    parse_contract,
    read_selected_contract,
)

TODAY = date(2026, 10, 10)


def workbook(count=1, change=None, duplicate=False):
    wb = Workbook()
    ws = wb.active
    ws["H1"] = "PI No.: SHAWN20260101-01"
    ws.append([])
    for col, text in enumerate(["PART NUMBER", "L/T", "BRAND", "CONDITION", "DC", "QTY(PCS)", "UNIT PRICE(￥)", "AMOUNT"], 1):
        ws.cell(6, col, text)
    for r in range(7, 7 + count):
        for c, value in enumerate([f"SYNTHETIC-{r}", "1-3 DAYS", "TEST", "unused", "24+", 2, 3.25, "=F7*G7"], 1):
            ws.cell(r, c, value)
    footer = 7 + count
    ws.merge_cells(start_row=footer, start_column=1, end_row=footer, end_column=7)
    ws.cell(footer, 1, "TOTAL PAYMENT")
    if change:
        for cell, value in change.items(): ws[cell] = value
    if duplicate:
        other = wb.copy_worksheet(ws)
        other["H1"] = None
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize("n", [1, 4])
def test_parse_all_required_fields_amount_ignored(n):
    order = parse_contract(workbook(n), today=TODAY)
    assert len(order.lines) == n and order.today == TODAY
    assert all(x.quantity == 2 and x.unit_price == Decimal("3.25") and x.delivery_date == "2026-10-17" for x in order.lines)
    assert "SHAWN" not in repr(order) and "SYNTHETIC" not in repr(order.lines)


@pytest.mark.parametrize("cell,value", [("A7", None), ("B7", ""), ("C7", None), ("E7", None),
    ("F7", 0), ("F7", -1), ("F7", 1.5), ("F7", "3"), ("F7", True),
    ("G7", 0), ("G7", -1), ("G7", "3"), ("G7", "=1+2"), ("E7", "=1"),
    ("B7", "soon"), ("H1", "PI No.: SHAWN20260101-01 SHAWN20260101-02")])
def test_invalid_contract_stops(cell, value):
    with pytest.raises(OrderError): parse_contract(workbook(change={cell: value}), today=TODAY)


def test_ambiguous_or_missing_header_stops():
    with pytest.raises(OrderError): parse_contract(workbook(duplicate=True), today=TODAY)
    with pytest.raises(OrderError): parse_contract(workbook(change={"E6": "missing"}), today=TODAY)


@pytest.mark.parametrize("lt,days", [("1-3 DAYS", 7), ("2-4 WEEKS", 35), (" 3-6 weeks ", 49)])
def test_lead_time_rules(lt, days):
    assert (date.fromisoformat(delivery_date(lt, TODAY)) - TODAY).days == days


@pytest.mark.parametrize("lt", ["1 DAYS", "2-3 DAYS", "4 WEEKS", "3-2 WEEKS", "0-2 WEEKS", "SOON"])
def test_unsupported_lead_time_stops(lt):
    with pytest.raises(OrderError, match="LEAD_TIME"): delivery_date(lt, TODAY)


def test_selected_mail_only_and_flags_required():
    calls = []
    def receiver(**kwargs):
        calls.append(kwargs["selected_date"])
        message = EmailMessage()
        message["Date"] = "Wed, 07 Oct 2026 12:26:16 +0800"
        message.set_content("synthetic")
        message.add_attachment(workbook(), maintype="application", subtype="octet-stream", filename="test.xlsx")
        kwargs["_message_consumer"](message)
        return SimpleNamespace(status="SUCCESS", flags_unchanged=True)
    assert len(read_selected_contract(2, today=TODAY, inspect=receiver).lines) == 1
    assert len(calls) == 1 and calls[0].startswith("2026-10-07")
