from email.message import EmailMessage
from io import BytesIO

from openpyxl import Workbook

from src.core import CredentialSiteNotFoundError, Login
from src.order_mail.inspection import describe_message, inspect_order_mail


def message(subject="订单录单 SYNTHETIC-123"):
    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = "synthetic@example.test"
    msg["Date"] = "Fri, 09 Oct 2026 16:45:52 +0800"
    msg["Message-ID"] = "<synthetic-id@example.test>"
    msg.set_content("SENSITIVE BODY")
    wb = Workbook()
    ws = wb.active
    ws.title = "SENSITIVE CUSTOMER"
    ws["A1"] = "合同号：SENSITIVE CONTRACT"
    ws["A2"] = "客户名称：SENSITIVE CUSTOMER"
    ws["B3"] = "数量"
    ws["B4"] = 12345
    ws["C4"] = "=B4*10"
    ws.merge_cells("A1:D1")
    hidden = wb.create_sheet("hidden private")
    hidden.sheet_state = "hidden"
    buf = BytesIO()
    wb.save(buf)
    msg.add_attachment(buf.getvalue(), maintype="application", subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename="SENSITIVE CONTRACT.xlsx")
    return msg


class Client:
    def __init__(self, *, subject="订单录单 SYNTHETIC-123", changed=False, readonly=True, large=False):
        self.msg = message(subject)
        self.commands = []
        self.changed, self.is_readonly, self.large = changed, readonly, large
        self.logged_out = False

    def login(self, username, password):
        assert username == "linan229@qq.com"
        assert password == "synthetic-secret"
        return "OK", [b"ok"]

    def select(self, mailbox, readonly):
        assert mailbox == "INBOX" and readonly
        return "OK", [b"250"]

    def uid(self, command, *args):
        self.commands.append((command, args))
        if command == "search":
            assert args[-1] == "51:250"
            assert args[0:2] == ("CHARSET", "UTF-8")
            return "OK", [b"1 2 3 4 5 6 7 8"]
        assert command == "fetch"
        uid, fields = args
        flag = b"\\Seen" if self.changed and fields == "(UID FLAGS)" else b""
        size = 30_000_000 if self.large else len(self.msg.as_bytes())
        meta = b'1 (UID ' + uid + b' FLAGS (' + flag + b') INTERNALDATE "09-Oct-2026 16:45:52 +0800" RFC822.SIZE ' + str(size).encode() + b')'
        if "PEEK" in fields:
            return "OK", [(meta, self.msg.as_bytes()), b")"]
        assert fields == "(UID FLAGS)"
        return "OK", [meta]

    def logout(self):
        self.logged_out = True


def run(client):
    def factory(host, port, **kwargs):
        assert (host, port) == ("imap.qq.com", 993)
        assert kwargs["ssl_context"].check_hostname
        assert kwargs["timeout"] == 20
        return client
    return inspect_order_mail(credential_getter=lambda _: Login("imap.qq.com", "", "linan229@qq.com", "synthetic-secret"), client_factory=factory)


def test_peek_five_candidates_and_no_sensitive_output():
    client = Client()
    report = run(client)
    assert report.status == "SUCCESS" and report.matched == 5
    assert report.flags_unchanged and client.logged_out
    assert "SENSITIVE" not in report.text and "SYNTHETIC" not in report.text
    assert "12345" not in report.text and "synthetic-id" not in report.text
    assert "合同号@A1" in report.text and "客户名称@A2" in report.text
    assert "hidden" in report.text and "公式 1" in report.text
    assert {args[0] for cmd, args in client.commands if cmd == "fetch"} == {b"4", b"5", b"6", b"7", b"8"}
    assert all("PEEK" in args[1] or args[1] == "(UID FLAGS)" for cmd, args in client.commands if cmd == "fetch")


def test_substring_subject_is_not_read_as_order():
    client = Client(subject="Re: 订单录单 SENSITIVE")
    report = run(client)
    assert report.matched == 0 and report.status == "SUCCESS"
    assert not any(args[1] == "(UID FLAGS BODY.PEEK[])" for cmd, args in client.commands if cmd == "fetch")


def test_flag_change_never_claims_unchanged_or_writes():
    client = Client(changed=True)
    report = run(client)
    assert report.status == "FLAGS_DIFFER" and report.flags_unchanged is False
    assert {cmd for cmd, _ in client.commands} == {"search", "fetch"}


def test_not_readonly_fails_closed_before_search():
    client = Client(readonly=False)
    report = run(client)
    assert report.status == "CHECK_FAILED" and not client.commands and client.logged_out


def test_large_messages_never_fetch_body():
    client = Client(large=True)
    report = run(client)
    assert report.matched == 5 and "20MB" in report.text
    assert not any(args[1] == "(UID FLAGS BODY.PEEK[])" for cmd, args in client.commands if cmd == "fetch")


def test_missing_credentials_no_network():
    def missing(_):
        raise CredentialSiteNotFoundError("imap.qq.com", "not present")
    def forbidden(*args, **kwargs):
        raise AssertionError("Network must not start")
    report = inspect_order_mail(credential_getter=missing, client_factory=forbidden)
    assert report.status == "CREDENTIAL_REQUIRED" and not report.connected


def test_wrong_account_no_network():
    report = inspect_order_mail(credential_getter=lambda _: Login("imap.qq.com", "", "wrong@example.test", "secret"), client_factory=lambda *a, **kw: (_ for _ in ()).throw(AssertionError()))
    assert report.status == "ACCOUNT_MISMATCH"


def test_untrusted_filename_mime_errors_and_body_never_escape():
    msg = message()
    msg.add_attachment(b"broken", maintype="application", subtype="sensitive-type", filename="SENSITIVE.pdf")
    lines = "\n".join(describe_message(msg, b""))
    assert "SENSITIVE" not in lines and "sensitive-type" not in lines
    assert "文件签名异常" in lines
