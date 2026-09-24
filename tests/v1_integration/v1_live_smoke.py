"""V1 final live integration smoke.

Runs the full V1 mainline with the real Google Sheet and the production
research runtime. Source adapters run for real; CAPTCHA / OTP / manual
login keep failing closed at the source layer (no bypass).

Required (Owner-supplied, never recorded in the repo):

    SPREADSHEET_ID       — Google Sheets spreadsheet ID
    WORKSHEET_TITLE      — single worksheet name, e.g. "2026"
    CLIENT_SECRET_FILE   — OAuth Desktop client secret JSON path

Optional env overrides:

    WORKSHEET_TITLES     — comma-separated list, e.g. "2026,shahab"
    INSO_SITE_ID=yingsuo.alperp.cn
    BRAND_UPDATER=noop   — always disabled for V1 smoke (no write-back)

Run from the worktree root:

    SPREADSHEET_ID=... WORKSHEET_TITLES=2026,shahab CLIENT_SECRET_FILE=... \\
        python tests/v1_integration/v1_live_smoke.py
"""

from __future__ import annotations

import os
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from src.research.runtime import (
    assess_readiness,
    build_production_research_service,
    format_readiness,
    load_runtime_config,
)
from src.sheets import WorksheetIdentity
from src.sheets.google_oauth import (
    build_read_only_google_sheets_service,
)
from src.sheets.google_reader import GoogleSheetsRowReader
from src.workflow import (
    WorkflowPoller,
    WorkflowRuntime,
    WorkflowStateStore,
    WorkflowWorker,
)


def _require_env(name: str) -> str:
    value = os.environ.get(name, "").strip()
    if not value:
        raise SystemExit(f"Missing required env var: {name}")
    return value


def _worksheet_titles() -> tuple[str, ...]:
    multi = os.environ.get("WORKSHEET_TITLES", "").strip()
    if multi:
        return tuple(title.strip() for title in multi.split(",") if title.strip())
    single = os.environ.get("WORKSHEET_TITLE", "").strip()
    if single:
        return (single,)
    raise SystemExit("Missing required env var: WORKSHEET_TITLES or WORKSHEET_TITLE")


def main() -> int:
    spreadsheet_id = _require_env("SPREADSHEET_ID")
    worksheet_titles = _worksheet_titles()
    client_secret = Path(_require_env("CLIENT_SECRET_FILE"))
    if not client_secret.exists():
        raise SystemExit(f"client_secret file not found: {client_secret}")

    config = load_runtime_config()
    runtime_dir = ROOT / "runtime"
    runtime_dir.mkdir(parents=True, exist_ok=True)
    db_path = runtime_dir / "v1_live.sqlite"
    if db_path.exists():
        raise SystemExit(f"Refusing to overwrite existing smoke database: {db_path}")

    readiness = assess_readiness(config.cdp.cdp_url)
    print(format_readiness(readiness))
    if not readiness.ready:
        raise SystemExit("runtime not ready; fix blockers above")

    service = build_read_only_google_sheets_service(client_secret)
    reader = GoogleSheetsRowReader(service)
    worksheets = [
        WorksheetIdentity(spreadsheet_id, title) for title in worksheet_titles
    ]

    excel_path = config.excel_output_path
    if excel_path.exists():
        raise SystemExit(f"Refusing to overwrite existing Excel output: {excel_path}")

    store = WorkflowStateStore(db_path)
    poller = WorkflowPoller(store, reader)
    research = build_production_research_service(
        config,
        cdp_probe=lambda url: True,
    )
    # Brand updater is disabled: we record NOT_CONFIGURED instead of writing.
    worker = WorkflowWorker(store, research, brand_updater=None)
    runtime = WorkflowRuntime(poller, worker, worksheets)

    now = datetime.now(timezone.utc).replace(microsecond=0)
    print("== first poll ==")
    added = runtime.run_poll(now=now)
    print(f"  added = {added}")

    print("== second poll (no new pending) ==")
    added2 = runtime.run_poll(now=now)
    print(f"  added = {added2}")
    assert added2 == 0, f"second poll added {added2} duplicates"

    print("== Research worker drain ==")
    processed = runtime.drain_due(now=now)
    print(f"  processed = {processed}")

    # Excel summary
    excel_path = config.excel_output_path
    if excel_path.exists():
        from openpyxl import load_workbook

        wb = load_workbook(excel_path)
        ws_excel = wb.active
        print(f"  excel rows = {ws_excel.max_row}")
    else:
        print(f"  excel not present at {excel_path}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
